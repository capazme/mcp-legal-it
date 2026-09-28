"""Live gate (strutturale): `genera_modello_atto` and `esporta_atto_docx`.

Phase 4 of the avvocatoandreani.it benchmark, strategy `strutturale`, group `modelli_atti`.
The benchmark site has no counterpart for either tool, so this file checks what they produce
against the vigente text of the norms (Normattiva / EUR-Lex through `cite_law`) and against
the registered MCP surface:

* `genera_modello_atto` is a catalogue (src/data/modelli_atti.json, 100 entries): every entry
  must route to a registered tool with parameters that tool accepts, its `riferimenti_normativi`
  must exist (`verifica_citazioni`) and be in force (no "ARTICOLO/PROVVEDIMENTO ABROGATO"), and
  its `avvertenze` must match the vigente text of the article they cite. The plan's four cases
  are here: attestazione di copia informatica (artt. 196-octies ss. disp. att. c.p.c. in place of
  the abrogated artt. 16-bis co. 9-bis and 16-undecies DL 179/2012), search for "196-octies",
  atto di precetto (artt. 479-481 c.p.c., art. 480 co. 2 warning in the direct tool) and
  citazione ordinaria (art. 163-bis c.p.c., 120/150 days).
* `esporta_atto_docx` writes a .docx under /tmp/mcp-legal-it: each produced file is copied into
  a `tempfile` directory, deleted from /tmp/mcp-legal-it straight away, and opened with
  python-docx to check headings, runs, list styles, core properties, numbering and the
  citations it carries (again through `verifica_citazioni`).

A failing test here is a genuine divergence of the catalogue or of the converter from the norm
or from the tool's own contract: do not soften it. It needs the network and is skipped by
default:

    .venv/bin/pytest tests/unit/test_strutturale_modelli_atti_live.py -m live -q -p no:cacheprovider -rfEs
"""

from __future__ import annotations

import asyncio
import functools
import json
import os
import re
import shutil
import tempfile
from pathlib import Path

import pytest

from tests.unit._norme_live import assert_parole, cite_law_json, contiene, normalizza

pytestmark = pytest.mark.live

_REPO = Path(__file__).resolve().parents[2]
_CATALOGO_JSON = _REPO / "src" / "data" / "modelli_atti.json"
_DOCX_DIR = "/tmp/mcp-legal-it"  # fixed by the tool (src/tools/modelli_atti.py)
_NORMATTIVA = "https://www.normattiva.it/uri-res/N2Ls?urn:nir:stato:"
_DISP_ATT_CPC = "regio.decreto:1941-08-25;1368:1"
_BATCH = 20  # verifica_citazioni resolves at most 20 references per call


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _tool(module: str, name: str):
    import importlib

    import src.server  # noqa: F401  (registers every tool module first)

    obj = getattr(importlib.import_module(f"src.tools.{module}"), name)
    return getattr(obj, "fn", obj)


def _genera(tipo_atto: str, parametri: dict | None = None) -> dict:
    return _tool("modelli_atti", "genera_modello_atto")(tipo_atto=tipo_atto, parametri=parametri)


@functools.lru_cache(maxsize=1)
def _catalogo() -> dict[str, dict]:
    raw = json.loads(_CATALOGO_JSON.read_text(encoding="utf-8"))
    return {k: v for k, v in raw.items() if not k.startswith("_")}


@functools.lru_cache(maxsize=1)
def _vintage() -> dict:
    return json.loads(_CATALOGO_JSON.read_text(encoding="utf-8")).get("_vintage", {})


def _entry_text(tipo_atto: str) -> str:
    """Everything the catalogue says about an entry, as one normalised string."""
    e = _catalogo()[tipo_atto]
    parti = [e["descrizione"], *e.get("riferimenti_normativi", []), *e.get("avvertenze", [])]
    return normalizza(" | ".join(parti))


@functools.lru_cache(maxsize=1)
def _superficie_mcp() -> tuple[dict[str, dict], list[str]]:
    """Registered tools (name -> inputSchema) and static resources, read in-process."""
    from fastmcp import Client

    from src.server import mcp

    async def run():
        async with Client(mcp) as client:
            tools = {t.name: t.inputSchema for t in await client.list_tools()}
            resources = [str(r.uri) for r in await client.list_resources()]
            return tools, resources

    return asyncio.run(run())


def _verifica(citazioni: list[str]) -> list[dict]:
    """Run `verifica_citazioni` (JSON) over the references, in batches of 20."""
    fn = _tool("legal_citations", "verifica_citazioni")
    out: list[dict] = []
    for i in range(0, len(citazioni), _BATCH):
        payload = json.loads(asyncio.run(fn("\n".join(citazioni[i:i + _BATCH]), formato="json")))
        assert payload.get("errore") is None, payload
        assert not payload.get("troncato"), payload
        out.extend(payload["citazioni"])
    return out


_ART = r"\d+(?:-[a-z]+)?"
_SOFT_LAW = re.compile(r"linee guida|wp\s?\d+|specifiche tecniche", re.IGNORECASE)


def _normalizza_riferimento(ref: str) -> list[str]:
    """Turn a catalogue reference into citations `cite_law` / `verifica_citazioni` can resolve.

    Ranges ("artt. 633-656 c.p.c.") are checked at both ends, "artt. X ss." at X, whole acts
    through their art. 1 (proves the act exists); soft law (EDPB/WP29 guidelines, DGSIA
    technical specifications) is not a norm and is skipped. A bare "art. 3-bis" with no act is
    returned as is: it is a defect of the catalogue and `verifica_citazioni` reports it.
    """
    r = ref.strip()
    if _SOFT_LAW.search(r):
        return []
    m = re.match(r"^D\.P\.C\.M\.\s+\d{1,2}/\d{1,2}/(\d{4})\s+n\.\s*(\d+)$", r)
    if m:  # "D.P.C.M. 16/2/2016 n. 40": the resolver wants number/year
        return [f"art. 1 DPCM {m[2]}/{m[1]}"]
    m = re.match(r"^(?P<atto>.+?)\s+(?P<k>artt?\.)\s+(?P<rest>.+)$", r)
    if m and not r.lower().startswith("art"):  # "DPR 115/2002 art. 13 co. 3": act first
        r = f"{m['k']} {m['rest']} {m['atto']}"
    r = re.sub(r"\s*\(GDPR\)", "", r)
    r = re.sub(r"\s+conv\. L\. \d+/\d{4}", "", r)
    r = re.sub(r"\s+agg\. DM \d+/\d{4}.*$", "", r)
    r = r.replace(" co. 9-bis", "")  # "art. 16-bis co. 9-bis DL ...": the parser reads "-bis DL" as the act
    m = re.match(rf"^artt\.\s+({_ART})\s*-\s*({_ART})\s+(.+)$", r)
    if m:
        return [f"art. {m[1]} {m[3]}", f"art. {m[2]} {m[3]}"]
    m = re.match(rf"^artt\.\s+({_ART})\s+ss\.\s+(.+)$", r)
    if m:
        return [f"art. {m[1]} {m[2]}"]
    if r.lower().startswith("art."):
        return [r]
    if re.search(r"\d+/\d{4}", r):
        return [f"art. 1 {r}"]
    return []


@functools.lru_cache(maxsize=1)
def _citazioni_catalogo() -> dict[str, list[str]]:
    """Normalised citation -> catalogue entries that cite it (through the original reference)."""
    mappa: dict[str, list[str]] = {}
    for tipo, e in _catalogo().items():
        for ref in e.get("riferimenti_normativi", []):
            for cit in _normalizza_riferimento(ref):
                mappa.setdefault(cit, []).append(tipo)
    return mappa


_ABROGATO = re.compile(r"(?:ARTICOLO|PROVVEDIMENTO) ABROGATO")


def _abrogato(reference: str) -> str | None:
    """The abrogation notice at the head of the vigente text, if any.

    One retry after a pause: CELLAR occasionally answers an EU article with an empty body
    ("nessun testo trovato") and serves it on the next request.
    """
    import time

    payload = cite_law_json(reference)
    if payload.get("errore"):
        time.sleep(3)
        payload = cite_law_json(reference)
    assert not payload.get("errore"), payload
    head = (payload.get("testo") or "")[:600]
    m = _ABROGATO.search(head)
    return head[m.start():m.start() + 120].splitlines()[0] if m else None


# ---------------------------------------------------------------------------
# genera_modello_atto — the plan's cases
# ---------------------------------------------------------------------------

def test_piano_attestazione_copia_informatica_cita_artt_196_disp_att_cpc():
    """Plan case 1: the entry must cite the vigente PCT norms (artt. 196-novies/196-decies and
    196-undecies disp. att. c.p.c., D.Lgs. 149/2022), not only art. 16-bis co. 9-bis DL 179/2012,
    which art. 11 D.Lgs. 149/2022 abrogated."""
    assert_parole("art. 11 D.Lgs. 149/2022", "16-bis", "16-undecies", "sono abrogati")
    assert _abrogato("art. 16-bis DL 179/2012"), "art. 16-bis DL 179/2012 no longer marked abrogated"
    assert_parole("art. 196-novies disp. att. c.p.c.", "formato su supporto analogico",
                  "attestano la conformita' della copia")
    assert_parole("art. 196-decies disp. att. c.p.c.", "all'ufficiale giudiziario",
                  "formato su supporto analogico")
    assert_parole("art. 196-undecies disp. att. c.p.c.", "modalita' dell'attestazione di conformita'")

    r = _genera("attestazione_copia_informatica")
    assert r["tool_diretto"] == "attestazione_conformita"
    assert r["parametri_fissi"] == {"modalita": "copia_informatica"}
    rif = normalizza(" ".join(r["riferimenti_normativi"]))
    assert "196-undecies" in rif and ("196-novies" in rif or "196-decies" in rif), (
        f"attestazione_copia_informatica cita {r['riferimenti_normativi']}: norme abrogate dall'art. 11 "
        "D.Lgs. 149/2022 (applicabili solo ai procedimenti pendenti al 28/02/2023, art. 35 co. 1); "
        "servono gli artt. 196-novies (deposito) / 196-decies (trasmissione all'UG) e 196-undecies "
        "disp. att. c.p.c."
    )


def test_piano_cerca_196_octies_trova_le_attestazioni():
    """Plan case 2: searching the catalogue for the norm that now governs attestations of
    conformity (art. 196-octies disp. att. c.p.c.) must return the attestazioni entries."""
    r = _genera("cerca", {"query": "196-octies"})
    assert "errore" not in r, r
    attestazioni = sorted(k for k, v in _catalogo().items() if v["categoria"] == "attestazioni")
    trovati = sorted(x["tipo_atto"] for x in r["risultati"])
    assert r["totale"] >= 1 and set(trovati) & set(attestazioni), (
        f"cerca '196-octies' -> {r['totale']} risultati; attese le voci {attestazioni}"
    )


def test_piano_atto_di_precetto_artt_479_481_e_avvertimento_art_480_co_2():
    """Plan case 3: artt. 479-481 c.p.c.; 90-day inefficacy (art. 481 co. 1); the direct tool's
    text carries the art. 480 co. 2 warning on the sovraindebitamento remedies."""
    assert_parole("art. 481 c.p.c.", "novanta giorni dalla sua notificazione")
    assert_parole("art. 480 c.p.c.", "sovraindebitamento", "piano del consumatore",
                  "termine non minore di dieci giorni")

    r = _genera("atto_di_precetto")
    assert r["riferimenti_normativi"] == ["artt. 479-481 c.p.c."]
    assert r["tool_diretto"] == "atto_di_precetto"
    assert not contiene(" ".join(r["avvertenze"]), "90 giorni", "art. 481")

    precetto = _tool("atti_giudiziari", "atto_di_precetto")(
        creditore="Alfa S.r.l.", debitore="Mario Rossi",
        titolo_esecutivo="sentenza Tribunale di Milano n. 1234/2025", importo_capitale=10000.0,
    )
    missing = contiene(precetto["testo"], "art. 480 co. 2", "sovraindebitamento",
                       "piano del consumatore", "dieci giorni")
    assert not missing, f"testo del precetto privo di {missing}"


def test_piano_citazione_ordinaria_art_163_bis_120_150_giorni():
    """Plan case 4: art. 163 c.p.c. and D.Lgs. 149/2022; termini a comparire of at least 120 days
    (150 abroad) under art. 163-bis co. 1 as rewritten by the Cartabia reform."""
    assert_parole("art. 163-bis c.p.c.", "centoventi giorni", "centocinquanta giorni")
    r = _genera("citazione_ordinaria")
    assert r["riferimenti_normativi"] == ["art. 163 c.p.c.", "D.Lgs. 149/2022"]
    assert not contiene(" ".join(r["avvertenze"]), "120 giorni", "150 giorni")
    assert r["resource_modello"] == "legal://modelli-atti/introduttivi"
    assert "non ancora disponibile" in r["istruzioni"]


# ---------------------------------------------------------------------------
# genera_modello_atto — attestazioni: which norm governs which proceeding
# ---------------------------------------------------------------------------

def _albero_disp_att() -> str:
    """Normalised table of contents of the disp. att. c.p.c. as Normattiva prints it today."""
    import httpx
    from bs4 import BeautifulSoup

    from src.lib import _clock
    from src.lib.visualex.scraper import _HEADERS

    url = f"{_NORMATTIVA}{_DISP_ATT_CPC}~art196octies!vig={_clock.today().isoformat()}"
    with httpx.Client(headers=_HEADERS, timeout=40, follow_redirects=True) as client:
        resp = client.get(url)
        resp.raise_for_status()
    testo = BeautifulSoup(resp.text, "lxml").get_text(" ", strip=True)
    return normalizza(testo.replace("((", "").replace("))", ""))


def test_regime_transitorio_attestazioni_art_35_dlgs_149_2022():
    """Grounds the transitional statement used in the findings: art. 35 co. 2 D.Lgs. 149/2022
    extends to pending proceedings only capo I of titolo V-ter (artt. 196-quater - 196-septies)
    and art. 196-duodecies; the attestations (capo II, artt. 196-octies - 196-undecies) follow
    co. 1, so artt. 16-bis co. 9-bis and 16-undecies DL 179/2012 still govern the proceedings
    pending on 28/02/2023."""
    testo = assert_parole(
        "art. 35 D.Lgs. 149/2022",
        "si applicano ai procedimenti instaurati successivamente a tale data",
        "capo i del titolo vter|capo i del titolo v-ter",
    )
    assert "capo ii del titolo v" not in testo
    albero = _albero_disp_att()
    assert re.search(
        r"capo i degli atti e dei provvedimenti art\. 196 quater art\. 196 quinquies "
        r"art\. 196 sexies art\. 196 septies capo ii della conformit\S* delle copie agli originali "
        r"art\. 196 octies art\. 196 novies art\. 196 decies art\. 196 undecies capo iii",
        albero,
    ), "struttura del titolo V-ter disp. att. c.p.c. cambiata"


def test_attestazioni_non_citano_solo_norme_abrogate():
    """All 11 attestazioni entries: each must cite at least one art. 196-* disp. att. c.p.c.;
    artt. 16-bis co. 9-bis / 16-undecies DL 179/2012 may appear only as the regime of the
    proceedings pending on 28/02/2023."""
    assert _abrogato("art. 16-undecies DL 179/2012"), "art. 16-undecies no longer marked abrogated"
    solo_abrogate = {
        k: v["riferimenti_normativi"]
        for k, v in _catalogo().items()
        if v["categoria"] == "attestazioni"
        and not any("196-" in r for r in v["riferimenti_normativi"])
    }
    assert not solo_abrogate, (
        f"{len(solo_abrogate)} voci di attestazione citano solo norme abrogate dall'art. 11 "
        f"D.Lgs. 149/2022: {solo_abrogate}"
    )


# ---------------------------------------------------------------------------
# genera_modello_atto — catalogue structure against the registered MCP surface
# ---------------------------------------------------------------------------

def test_catalogo_100_tipi_10_categorie_e_footer_dati_applicati():
    r = _genera("catalogo")
    assert r["totale_tipi"] == 100 == len(_catalogo())
    assert len(r["categorie"]) == 10
    assert sum(len(v) for v in r["catalogo"].values()) == 100
    # @sourced footer: the catalogue declares no validity window (see _vintage in the result).
    assert any("modelli atti" in d for d in r["dati_applicati"]), r["dati_applicati"]
    assert _vintage().get("verifica") == "manuale"


def test_routing_e_tool_calcolo_puntano_a_tool_registrati():
    tools, _ = _superficie_mcp()
    mancanti = []
    for k, e in _catalogo().items():
        routing = e["routing"]
        if routing.get("tool") and routing["tool"] not in tools:
            mancanti.append((k, routing["tool"]))
        mancanti += [(k, c) for c in e.get("tool_calcolo", []) if c not in tools]
    assert not mancanti, mancanti


def test_resource_e_preventivo_procedura_dichiarati_non_disponibili():
    """47 entries have no generator yet (30 resource entries, 17 preventivo_procedura): the tool
    must say so instead of pointing to a resource or tool that does not exist."""
    tools, resources = _superficie_mcp()
    assert "preventivo_procedura" not in tools
    for k, e in _catalogo().items():
        tipo = e["routing"]["tipo"]
        if tipo == "resource":
            assert e["routing"]["resource"] not in resources
            assert "non ancora disponibile" in _genera(k)["istruzioni"], k
        elif tipo == "preventivo_procedura":
            r = _genera(k)
            assert r["tool_diretto"] == "preventivo_civile" and "approssimazione" in r["istruzioni"], k


def test_parametri_fissi_accettati_dal_tool_di_destinazione():
    """`parametri_fissi` are returned with "Chiamare/Usare il tool X": every key must be a
    parameter of X, otherwise the call fails with 'Unexpected keyword argument'."""
    tools, _ = _superficie_mcp()
    rifiutati = []
    for k, e in _catalogo().items():
        routing = e["routing"]
        if not routing.get("tool"):
            continue
        props = tools[routing["tool"]].get("properties", {})
        rifiutati += [
            f"{k}: {routing['tool']}({p}=...)"
            for p in (routing.get("parametri_fissi") or {})
            if p not in props
        ]
    assert not rifiutati, f"{len(rifiutati)} voci passano parametri inesistenti: {rifiutati}"


def test_campi_obbligatori_coincidono_con_i_parametri_del_tool_diretto():
    """For `tool_diretto` entries the catalogue's `campi_obbligatori` drive `campi_mancanti`:
    they must be parameters of the direct tool, and the tool's required parameters must be
    among them (or fixed), or `campi_mancanti == []` does not mean the call will succeed."""
    tools, _ = _superficie_mcp()
    difformi = {}
    for k, e in _catalogo().items():
        routing = e["routing"]
        if routing["tipo"] != "tool_diretto":
            continue
        schema = tools[routing["tool"]]
        props, required = set(schema.get("properties", {})), set(schema.get("required", []))
        campi = set(e["campi_obbligatori"]) | set(e.get("campi_opzionali", []))
        estranei = sorted(campi - props)
        scoperti = sorted(required - set(e["campi_obbligatori"]) - set(routing.get("parametri_fissi") or {}))
        if estranei or scoperti:
            difformi[k] = {"non_parametri": estranei, "obbligatori_del_tool_mancanti": scoperti}
    assert not difformi, f"{len(difformi)} voci tool_diretto difformi dal tool: {difformi}"


# ---------------------------------------------------------------------------
# genera_modello_atto — riferimenti normativi: existence and vigenza
# ---------------------------------------------------------------------------

def test_riferimenti_normativi_esistono_verifica_citazioni():
    """Every reference in the catalogue, normalised, must be `verificata` by verifica_citazioni."""
    citazioni = sorted(_citazioni_catalogo())
    assert len(citazioni) > 80
    esiti = _verifica(citazioni)
    non_verificate = {
        c["citazione"]: (c["verdetto"], _citazioni_catalogo()[c["citazione"]])
        for c in esiti
        if c["verdetto"] != "verificata"
    }
    assert not non_verificate, f"riferimenti non verificati: {non_verificate}"


def test_riferimenti_normativi_in_vigore():
    """No catalogue reference may point to an article or act Normattiva marks as abrogated.
    (verifica_citazioni answers `verificata` for an abrogated article: this test reads the text.)"""
    abrogati = {}
    for cit, voci in sorted(_citazioni_catalogo().items()):
        if not cit.lower().startswith("art. ") or not re.match(r"art\. \S+ \S", cit):
            continue
        nota = _abrogato(cit)
        if nota:
            abrogati[cit] = {"nota": nota, "voci": sorted(set(voci))}
    assert not abrogati, f"{len(abrogati)} riferimenti abrogati nel catalogo: {abrogati}"


# ---------------------------------------------------------------------------
# genera_modello_atto — avvertenze against the vigente text
# ---------------------------------------------------------------------------

# (tipo_atto, what the catalogue says, norm, words of the vigente text that support it)
_AVVERTENZE_CORRETTE = [
    ("citazione_ordinaria", ["120 giorni", "150 giorni"], "art. 163-bis c.p.c.",
     ["centoventi giorni", "centocinquanta giorni"]),
    ("opposizione_decreto_ingiuntivo", ["40 giorni"], "art. 641 c.p.c.", ["quaranta giorni"]),
    ("relata_pec_appello", ["30 gg"], "art. 325 c.p.c.", ["trenta giorni"]),
    ("relata_pec_appello", ["6 mesi"], "art. 327 c.p.c.", ["sei mesi"]),
    ("atto_appello", ["10 gg dalla notifica"], "art. 165 c.p.c.", ["entro dieci giorni dalla notificazione"]),
    ("atto_appello", ["aumentato della metà"], "art. 13 DPR 115/2002",
     ["aumentato della meta' per i giudizi di impugnazione"]),
    ("preventivo_decreto_ingiuntivo", ["cu dimezzato", "art. 13 co. 3"], "art. 13 DPR 115/2002",
     ["ridotto alla meta' per i processi speciali previsti nel libro iv, titolo i"]),
    ("decreto_ingiuntivo_retribuzioni", ["tre volte", "art. 9 co. 1-bis"], "art. 9 DPR 115/2002",
     ["tre volte l'importo previsto dall'articolo 76"]),
    ("atto_di_precetto", ["90 giorni", "art. 481"], "art. 481 c.p.c.", ["novanta giorni"]),
    ("atto_di_precetto", ["60 giorni", "90 se all'estero", "art. 644"], "art. 644 c.p.c.",
     ["sessanta giorni", "novanta giorni"]),
    ("termine_efficacia_titolo", ["90 giorni"], "art. 481 c.p.c.", ["novanta giorni"]),
    ("invito_negoziazione", ["30 giorni"], "art. 4 DL 132/2014", ["trenta giorni dalla ricezione"]),
    ("procura_negoziazione", ["circolazione"], "art. 3 DL 132/2014", ["danno da circolazione di veicoli"]),
    ("proroga_567_cpc", ["una sola volta"], "art. 567 c.p.c.", ["una sola volta"]),
    ("ricorso_giudice_pace", ["10.000"], "art. 7 c.p.c.", ["diecimila euro"]),
    ("lettera_adeguamento_istat", ["75%"], "art. 32 L. 392/1978", ["75 per cento"]),
    ("sollecito_pagamento", ["bce + 8"], "art. 2 D.Lgs. 231/2002", ["maggiorato di otto punti percentuali"]),
    ("notifica_data_breach", ["72 ore"], "art. 33 GDPR", ["72 ore"]),
    ("relata_posta", ["consiglio dell'ordine"], "art. 1 L. 53/1994", ["autorizzazione del consiglio dell'ordine"]),
    ("relata_posta", ["registro cronologico"], "art. 8 L. 53/1994",
     ["registro cronologico", "non si applicano alle notifiche effettuate a mezzo posta elettronica certificata"]),
    ("richiesta_nominativi_morosi", ["morosi"], "art. 63 disp. att. c.c.", ["condomini morosi"]),
]


@pytest.mark.parametrize(
    ("tipo_atto", "nel_catalogo", "norma", "nel_testo"),
    _AVVERTENZE_CORRETTE,
    ids=[f"{t}-{n}" for t, _, n, _ in _AVVERTENZE_CORRETTE],
)
def test_avvertenze_coerenti_con_il_testo_vigente(tipo_atto, nel_catalogo, norma, nel_testo):
    assert not contiene(_entry_text(tipo_atto), *nel_catalogo), (tipo_atto, nel_catalogo)
    assert_parole(norma, *nel_testo)


def test_atto_appello_costituzione_appellato_venti_giorni_art_347():
    """art. 347 co. 1 c.p.c. (as amended by D.Lgs. 164/2024): the other parties constitute in
    appeal 'almeno venti giorni prima dell'udienza', not 70 days (art. 166 applies to the
    first-instance convenuto only)."""
    assert_parole("art. 347 c.p.c.", "le altre parti si costituiscono in appello almeno venti giorni prima dell'udienza")
    testo = _entry_text("atto_appello")
    assert "70 gg" not in testo and ("20 gg" in testo or "venti giorni" in testo), (
        f"atto_appello: {_catalogo()['atto_appello']['avvertenze']}"
    )


def test_perdita_efficacia_pignoramento_art_497_non_riguarda_iscrizione_a_ruolo():
    """art. 497 c.p.c.: 45 days without a request for sale or assignment. The iscrizione a ruolo
    deadline is 15 days (artt. 518 co. 6, 557 co. 2) or 30 days (art. 543 co. 4) from delivery."""
    assert_parole("art. 497 c.p.c.", "quarantacinque giorni senza che sia stata richiesta l'assegnazione o la vendita")
    assert_parole("art. 518 c.p.c.", "entro quindici giorni dalla consegna, a pena di inefficacia del pignoramento")
    assert_parole("art. 543 c.p.c.", "entro trenta giorni dalla consegna, a pena di inefficacia del pignoramento")
    avv = normalizza(" ".join(_catalogo()["perdita_efficacia_pignoramento"]["avvertenze"]))
    assert not ("45 gg" in avv and "iscrizione a ruolo" in avv), (
        f"perdita_efficacia_pignoramento: '{avv}' attribuisce all'art. 497 il termine d'iscrizione a ruolo"
    )


def test_ricerca_beni_492_bis_istanza_all_ufficiale_giudiziario():
    """art. 492-bis co. 1 c.p.c. (Cartabia): the creditor with titolo and precetto asks the
    ufficiale giudiziario; the presidente del tribunale authorises only before the precetto /
    art. 482 term when there is 'pericolo nel ritardo' (co. 2)."""
    assert_parole("art. 492-bis c.p.c.",
                  "su istanza del creditore munito del titolo esecutivo e del precetto, l'ufficiale giudiziario",
                  "se vi e' pericolo nel ritardo, il presidente del tribunale")
    avv = normalizza(" ".join(_catalogo()["ricerca_beni_492bis"]["avvertenze"]))
    assert not ("autorizzazione del presidente del tribunale" in avv and "pericolo" not in avv), (
        f"ricerca_beni_492bis: '{avv}' presenta come sempre necessaria l'autorizzazione presidenziale"
    )


def test_relata_pec_penale_non_si_fonda_sulla_l_53_1994():
    """art. 1 L. 53/1994 limits lawyers' notifications to civil, administrative and
    stragiudiziale matters; in criminal proceedings the defence's PEC notification is art. 152
    c.p.p."""
    assert_parole("art. 1 L. 53/1994", "in materia civile, amministrativa e stragiudiziale")
    assert_parole("art. 152 c.p.p.", "eseguita dal difensore a mezzo di posta elettronica certificata")
    rif = _catalogo()["relata_pec_penale"]["riferimenti_normativi"]
    assert any("152 c.p.p." in r for r in rif) and "L. 53/1994" not in rif, f"relata_pec_penale: {rif}"


def test_formula_esecutiva_abolita_art_475_cpc():
    """art. 475 c.p.c. (D.Lgs. 149/2022; D.Lgs. 164/2024 art. 7 co. 4 for titles put into execution
    after 28/02/2023): the title is issued 'in copia attestata conforme all'originale o in
    duplicato informatico' — the formula esecutiva no longer exists."""
    testo = assert_parole("art. 475 c.p.c.", "copia attestata conforme all'originale o in duplicato informatico")
    assert "formula" not in testo.split("aggiornamento")[0]
    con_formula = sorted(k for k in _catalogo() if "formula esecutiva" in _entry_text(k))
    assert not con_formula, f"voci che richiedono ancora la formula esecutiva: {con_formula}"


def test_lettera_adeguamento_istat_riferimento_pertinente():
    """'L. 449/1997 art. 2 co. 1' is the transfer of State housing to municipalities: it has
    nothing to do with the ISTAT update of a rent."""
    assert_parole("art. 2 L. 449/1997", "trasferimento di alloggi ai comuni")
    rif = _catalogo()["lettera_adeguamento_istat"]["riferimenti_normativi"]
    assert "L. 449/1997 art. 2 co. 1" not in rif, f"lettera_adeguamento_istat: {rif}"


def test_decreto_ingiuntivo_professionale_parcella_art_636():
    """The professional's parcella with the Ordine's opinion is the written proof of art. 636
    c.p.c.; art. 642 co. 1 grants provisional execution on cambiale, assegno, atto pubblico."""
    assert_parole("art. 636 c.p.c.", "corredata dal parere della competente associazione professionale")
    assert_parole("art. 642 c.p.c.", "cambiale, assegno bancario, assegno circolare")
    e = _catalogo()["decreto_ingiuntivo_professionale"]
    assert any("636" in r for r in e["riferimenti_normativi"]), (
        f"decreto_ingiuntivo_professionale: {e['riferimenti_normativi']} / {e['avvertenze']}"
    )


def test_sfratto_morosita_nessuna_soglia_di_due_canoni():
    """art. 5 L. 392/1978: one canone unpaid 20 days after the due date is grave breach; the
    'due mensilita'' threshold concerns the oneri accessori only."""
    assert_parole("art. 5 L. 392/1978", "mancato pagamento del canone decorsi venti giorni",
                  "oneri accessori quando l'importo non pagato superi quello di due mensilita'")
    avv = normalizza(" ".join(_catalogo()["sfratto_morosita"]["avvertenze"]))
    assert ">= 2 canoni" not in avv, f"sfratto_morosita: '{avv}'"


# ---------------------------------------------------------------------------
# esporta_atto_docx
# ---------------------------------------------------------------------------

_PIANO_TESTO = (
    "# Atto di citazione\n## Fatto\nIl **sig. Rossi** ha *consegnato* la merce.\n"
    "- primo\n- secondo\n1. uno\n2. due\n> nota\n### Diritto\nArt. 1218 c.c."
)


@pytest.fixture
def esporta():
    """Call esporta_atto_docx, move the file into a tempfile dir, yield (message, Document, path)."""
    from docx import Document

    fn = _tool("modelli_atti", "esporta_atto_docx")
    tmp = tempfile.mkdtemp(prefix="strutturale_docx_")

    def _run(**kwargs):
        msg = fn(**kwargs)
        m = re.match(r"File salvato: (.+\.docx) \(\d+(?:\.\d+)? KB\)$", msg)
        assert m, msg
        originale = m.group(1)
        copia = os.path.join(tmp, os.path.basename(originale))
        shutil.copyfile(originale, copia)
        os.remove(originale)  # never leave the tool's output behind
        return msg, Document(copia), originale

    yield _run
    shutil.rmtree(tmp, ignore_errors=True)


def _runs(par) -> list[tuple[str, bool, bool]]:
    return [(r.text, bool(r.bold), bool(r.italic)) for r in par.runs if r.text]


def test_piano_conversione_completa(esporta):
    """Plan case 1: Heading 1/2/3, bold and italic runs, List Bullet / List Number, italic quote,
    core properties title/author, sanitised file name."""
    msg, doc, originale = esporta(testo=_PIANO_TESTO, titolo="Citazione Rossi/Bianchi", autore="Avv. X")
    assert os.path.dirname(originale) == _DOCX_DIR
    assert re.fullmatch(r"citazione_rossibianchi_[0-9a-f]{8}\.docx", os.path.basename(originale))
    assert doc.core_properties.title == "Citazione Rossi/Bianchi"
    assert doc.core_properties.author == "Avv. X"
    got = [(p.style.name, _runs(p)) for p in doc.paragraphs]
    assert got == [
        ("Heading 1", [("Atto di citazione", False, False)]),
        ("Heading 2", [("Fatto", False, False)]),
        ("Normal", [("Il ", False, False), ("sig. Rossi", True, False), (" ha ", False, False),
                    ("consegnato", False, True), (" la merce.", False, False)]),
        ("List Bullet", [("primo", False, False)]),
        ("List Bullet", [("secondo", False, False)]),
        ("List Number", [("uno", False, False)]),
        ("List Number", [("due", False, False)]),
        ("Normal", [("nota", False, True)]),
        ("Heading 3", [("Diritto", False, False)]),
        ("Normal", [("Art. 1218 c.c.", False, False)]),
    ], got


def test_piano_testo_vuoto_nessun_file():
    """Plan case 2: blank text -> error message, no file written."""
    fn = _tool("modelli_atti", "esporta_atto_docx")
    prima = set(os.listdir(_DOCX_DIR)) if os.path.isdir(_DOCX_DIR) else set()
    assert fn(testo="   ", titolo="Vuoto Strutturale Benchmark") == "Errore: il testo dell'atto è vuoto."
    dopo = set(os.listdir(_DOCX_DIR)) if os.path.isdir(_DOCX_DIR) else set()
    assert not [f for f in dopo - prima if f.startswith("vuoto_strutturale_benchmark_")]


def test_titolo_con_path_traversal_resta_nella_cartella(esporta):
    _, _, originale = esporta(testo="Testo", titolo="../../etc/Evil Atto")
    assert os.path.dirname(originale) == _DOCX_DIR
    assert re.fullmatch(r"etcevil_atto_[0-9a-f]{8}\.docx", os.path.basename(originale))


def test_citazioni_nel_docx_verificate(esporta):
    """The citations the exported act carries survive the conversion verbatim and exist."""
    _, doc, _ = esporta(testo=_PIANO_TESTO, titolo="Citazioni")
    corpo = "\n".join(p.text for p in doc.paragraphs)
    citazioni = re.findall(r"[Aa]rt\.\s*\d+(?:-[a-z]+)?\s+c\.c\.", corpo)
    assert citazioni == ["Art. 1218 c.c."]
    esiti = _verifica(citazioni)
    assert [c["verdetto"] for c in esiti] == ["verificata"], esiti


def _numeri_resi(doc) -> list[int]:
    """Numbers Word shows for the List Number paragraphs: one counter per numId, no restart
    unless a paragraph carries its own numbering (numPr) pointing to another list."""
    from docx.oxml.ns import qn

    stile = doc.styles["List Number"].element
    num_stile = stile.find(f".//{qn('w:numId')}").get(qn("w:val"))
    contatori: dict[str, int] = {}
    resi = []
    for p in doc.paragraphs:
        if p.style.name != "List Number":
            continue
        ppr = p._p.pPr
        own = ppr.find(f".//{qn('w:numId')}") if ppr is not None else None
        num_id = own.get(qn("w:val")) if own is not None else num_stile
        contatori[num_id] = contatori.get(num_id, 0) + 1
        resi.append(contatori[num_id])
    return resi


def test_liste_numerate_conservano_la_numerazione_del_testo(esporta):
    """Two numbered lists separated by a paragraph, then a list starting at 10: Markdown (and the
    lawyer who wrote '1.' again) expects 1, 2 / 1, 2 / 10. The converter drops the typed number
    and every 'List Number' paragraph shares the style's single list, so Word shows 1..5."""
    testo = "1. primo\n2. secondo\n\nParagrafo\n\n1. nuova lista\n2. seconda\n\n10. decimo"
    _, doc, _ = esporta(testo=testo, titolo="Liste")
    assert _numeri_resi(doc) == [1, 2, 1, 2, 10]


def test_grassetto_e_corsivo_nei_titoli(esporta):
    """'## **Fatto** e *diritto*': the declared markers must not reach the heading as literal
    asterisks."""
    _, doc, _ = esporta(testo="## **Fatto** e *diritto*", titolo="Titoli")
    heading = doc.paragraphs[0]
    assert heading.style.name == "Heading 2"
    assert "*" not in heading.text, heading.text


def test_grassetto_corsivo_combinato_non_altera_il_testo_adiacente(esporta):
    """'***x***' (bold italic) must not leak asterisks nor turn the following plain words italic."""
    _, doc, _ = esporta(testo="Testo con ***grassetto corsivo*** e resto normale", titolo="Marcatori")
    par = doc.paragraphs[0]
    assert "*" not in par.text, par.text
    assert not any(italic for text, _, italic in _runs(par) if "resto" in text or text.strip() == "e"), _runs(par)


def test_titolo_usato_come_intestazione(esporta):
    """Docstring: 'titolo: Titolo del documento (usato come nome file e intestazione)'. The title
    must appear in the document (page header or opening heading), not only in the metadata."""
    _, doc, _ = esporta(testo="Corpo dell'atto.", titolo="Parere pro veritate")
    intestazione = " ".join(p.text for p in doc.sections[0].header.paragraphs)
    apertura = doc.paragraphs[0].text if doc.paragraphs else ""
    assert "Parere pro veritate" in intestazione or "Parere pro veritate" in apertura, (
        f"header={intestazione!r} primo paragrafo={apertura!r}"
    )
