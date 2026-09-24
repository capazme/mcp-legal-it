"""Wave-1 exit gates (handoff §5, DESIGN §8).

B1 power declared and satisfied (and refused when not); A4/B3 validity
cards present and coherent with the scorer; B2 judge panel with stubs
(reproducible order, heterogeneous families, calibration thresholds); B5
paired statistics; C ceiling/discrimination on the extended bank and
malformed items rejected with a useful message. No network: judges and the
model runner are stubbed.
"""

import json
import random
import subprocess
from pathlib import Path

import pytest

from benchmarks.limes.analysis.compare import contamination_report, wave_analysis
from benchmarks.limes.analysis.power import PowerError, enforce, power_check
from benchmarks.limes.analysis.stats import (
    holm,
    mcnemar_exact_power,
    ni_required_pairs,
    paired_non_inferiority,
    paired_risk_difference_newcombe,
)
from benchmarks.limes.bank import generate_q
from benchmarks.limes.bank.schema.item import ValidationError, load_bank
from benchmarks.limes.bank.schema.validate import validate_wave_slice
from benchmarks.limes.cli import main as cli_main
from benchmarks.limes.demo.simulation import simulate
from benchmarks.limes.protocol import judges
from benchmarks.limes.protocol import validity as registry
from benchmarks.limes.protocol.citations import has_marker, is_known_marker
from benchmarks.limes.protocol.hedge import detect_orientation, explicit_verdict
from benchmarks.limes.protocol.rules import load_protocol
from benchmarks.limes.runner.executor import _extract_payload, _parse_json_output, build_argv
from benchmarks.limes.runner.waves import WaveError, freeze_guard, load_wave

REPO_ROOT = Path(__file__).resolve().parents[2]
LIMES = REPO_ROOT / "benchmarks" / "limes"
BANK_DIR = LIMES / "bank"


@pytest.fixture(scope="module")
def wave():
    return load_wave(LIMES / "waves" / "wave-1.yaml")


@pytest.fixture(scope="module")
def bank(wave):
    return load_bank(BANK_DIR, wave.bank_slices, wave.require_validity, wave.excluded)


@pytest.fixture(scope="module")
def protocol():
    return load_protocol(LIMES / "protocol" / "protocol.yaml")


def _scored(bank):
    return [i for i in bank.items if not i.paraphrase_of]


# --------------------------------------------------------------------- bank


class TestExtendedBank:
    def test_size_and_construct_targets(self, bank):
        from benchmarks.limes.analysis.scorecard import construct_key

        scored = _scored(bank)
        assert len(scored) >= 150
        counts = {}
        for item in scored:
            key = construct_key(item.construct)
            counts[key] = counts.get(key, 0) + 1
        targets = {"C": 60, "H": 25, "P": 20, "A": 15, "U": 15, "M": 10}
        for key, minimum in targets.items():
            assert counts.get(key, 0) >= minimum, (key, counts)

    def test_twin_families_per_construct(self, bank):
        per = {}
        for a, _b in bank.twin_families().values():
            per[a.construct] = per.get(a.construct, 0) + 1
        assert per.get("analogia", 0) >= 6
        assert per.get("revirement", 0) >= 6
        assert per.get("gerarchia", 0) + per.get("diritto_ue", 0) >= 6

    def test_twins_have_opposite_orientations(self, bank):
        # The pair is the internal-consistency probe: one answer for both
        # branches must fail one of them.
        for family, (a, b) in bank.twin_families().items():
            if family.startswith(("H-", "A-", "U-")):
                assert a.correct_orientation != b.correct_orientation, family

    def test_every_item_has_validity_and_type(self, bank):
        for item in bank.items:
            assert item.validity is not None, item.id
            assert item.reasoning_type, item.id

    def test_validity_coherent_with_scorer(self, bank):
        # B3: the scorer the scorecard applies == the one the construct
        # definition declares (scorer_for is derived from code paths).
        for item in bank.items:
            registry.check_item(item)
            definition = registry.CONSTRUCT_DEFS[item.validity.construct_def]
            q_kind = item.q_answer.kind if item.q_answer else None
            assert registry.scorer_for(item.construct, item.layer, q_kind) == definition.scorer

    def test_registry_dimensions_match_scorecard_keys(self):
        from benchmarks.limes.analysis.scorecard import construct_key

        for definition in registry.CONSTRUCT_DEFS.values():
            assert construct_key(definition.construct) == definition.dimension

    def test_no_expected_marker_leaks_from_the_prompt(self, bank):
        for item in bank.items:
            if item.stratum == "anchor":
                continue
            for marker in item.expected_markers:
                assert not has_marker(item.prompt, marker), (item.id, marker)

    def test_paraphrases_point_to_anchors(self, bank):
        by_id = {i.id: i for i in bank.items}
        probes = [i for i in bank.items if i.paraphrase_of]
        assert len(probes) == 10
        for probe in probes:
            assert by_id[probe.paraphrase_of].stratum == "anchor"

    def test_wave1_slice_excludes_wave0_private(self, bank):
        assert not any(i.id.startswith("SP-") or i.id.startswith("QP-") for i in bank.items)


class TestGenerator:
    def test_deterministic(self):
        assert generate_q.generate("wave-1", 20260922) == generate_q.generate("wave-1", 20260922)

    def test_committed_slice_in_sync(self, capsys):
        assert generate_q.main(["--wave", "wave-1", "--check"]) == 0

    def test_other_wave_other_items(self):
        a = generate_q.generate("wave-1", 20260922)
        b = generate_q.generate("wave-2", 20260922)
        assert [i["prompt"] for i in a] != [i["prompt"] for i in b]

    def test_no_deadline_on_a_non_working_day(self):
        from datetime import date

        for item in generate_q.generate("wave-1", 20260922):
            if item["q_answer"]["kind"] == "data":
                d, m, y = (int(x) for x in item["q_answer"]["value"].split("/"))
                assert not generate_q.is_non_working(date(y, m, d)), item["id"]

    def test_month_end_rule(self):
        from datetime import date

        assert generate_q.add_months(date(2024, 8, 31), 18) == date(2026, 2, 28)
        assert generate_q.add_months(date(2023, 1, 31), 1) == date(2023, 2, 28)
        assert generate_q.add_months(date(2021, 3, 15), 60) == date(2026, 3, 15)

    def test_legal_interest_split_by_year(self):
        from datetime import date
        from decimal import Decimal

        # 83.400 EUR, 02/03/2025 -> 16/05/2026: 304 gg al 2% + 136 gg all'1,6%.
        total, steps = generate_q._legal_interest(Decimal(83400), date(2025, 3, 2), date(2026, 5, 16))
        assert round(float(total), 2) == 1886.44
        assert len(steps) == 2

    def test_easter_monday(self):
        from datetime import date

        assert generate_q.is_non_working(date(2026, 4, 6))  # Pasquetta 2026
        assert not generate_q.is_non_working(date(2026, 4, 7))


# --------------------------------------------------------------- B1 power


class TestPower:
    def test_required_sizes(self):
        assert ni_required_pairs(0.10, 0.15) == 118
        assert ni_required_pairs(0.10, 0.20) == 157
        assert ni_required_pairs(0.10, 0.25) == 197

    def test_mcnemar_power_table_recomputed(self):
        # The handoff table overstated m=30/40 at 75/25 (0.89/0.95).
        assert round(mcnemar_exact_power(20, 0.75), 2) == 0.62
        assert round(mcnemar_exact_power(30, 0.75), 2) == 0.80
        assert round(mcnemar_exact_power(40, 0.75), 2) == 0.90
        assert round(mcnemar_exact_power(20, 0.85), 2) == 0.93

    def test_wave1_declared_power_satisfied(self, wave, bank):
        report = power_check(wave.analysis["power"], len(_scored(bank)))
        assert report["declared"] and report["satisfied"], report

    def test_underpowered_wave_refused(self, wave):
        with pytest.raises(PowerError, match="potenza insufficiente"):
            enforce(wave.analysis["power"], 27)

    def test_plan_exits_nonzero_when_underpowered(self, tmp_path, capsys):
        import shutil

        root = tmp_path / "limes"
        shutil.copytree(LIMES, root, ignore=shutil.ignore_patterns("results", "__pycache__"))
        wave_file = root / "waves" / "wave-1.yaml"
        wave_file.write_text(
            wave_file.read_text(encoding="utf-8").replace("expected_discordance: 0.20", "expected_discordance: 0.60")
            .replace("margin: 0.10\n    alpha_one_sided", "margin: 0.05\n    alpha_one_sided"),
            encoding="utf-8",
        )
        code = cli_main(["plan", "--root", str(root), "--wave", "wave-1"])
        assert code == 1
        assert "potenza INSUFFICIENTE" in capsys.readouterr().out


# --------------------------------------------------------------- B5 stats


class TestPairedStatistics:
    def test_paired_newcombe_coverage(self):
        rng = random.Random(7)
        n, pa, p10, p01 = 150, 0.6, 0.12, 0.08
        covered = 0
        runs = 1500
        for _ in range(runs):
            cells = [0, 0, 0, 0]
            for _ in range(n):
                u = rng.random()
                cells[0 if u < pa else 1 if u < pa + p10 else 2 if u < pa + p10 + p01 else 3] += 1
            _, lo, hi = paired_risk_difference_newcombe(*cells)
            covered += lo <= p10 - p01 <= hi
        assert 0.93 <= covered / runs <= 0.97

    def test_non_inferiority_decision(self):
        same = [True] * 120 + [False] * 40
        assert paired_non_inferiority(same, list(same), 0.10)["non_inferior"]
        worse = [False] * 40 + [True] * 120
        control = [True] * 160
        assert not paired_non_inferiority(worse, control, 0.10)["non_inferior"]

    def test_holm(self):
        out = holm({"a": 0.01, "b": 0.04, "c": 0.03}, 0.05)
        assert out["a"]["rejected"] and not out["b"]["rejected"] and not out["c"]["rejected"]
        assert out["a"]["p_holm"] == pytest.approx(0.03)
        assert out["b"]["p_holm"] == pytest.approx(0.06)

    def test_wave_analysis_shape(self, wave):
        items = [f"i{k}" for k in range(60)]
        cells = {}
        for config, rate in (("bare", 0.4), ("mcp-legalit-v2.14.0-full", 0.7),
                             ("plugin-legalit-v2.14.0-full", 0.7),
                             ("plugin-legalit-v3.0.0-beta.3-full", 0.7)):
            cells[("m", config)] = {i: {"passed": k < rate * 60} for k, i in enumerate(items)}
        report = wave_analysis(cells, wave.analysis)
        primary = report["models"]["m"]["primary"]
        assert set(primary) == set(wave.analysis["primary"]["treatments"])
        assert all(row["rejected"] for row in primary.values())
        assert report["models"]["m"]["non_inferiority"]["non_inferior"]

    def test_contamination_flag(self):
        verdicts = {"o1": {"passed": True}, "p1": {"passed": False},
                    "o2": {"passed": True}, "p2": {"passed": False}}
        report = contamination_report(verdicts, [("o1", "p1"), ("o2", "p2")], 0.20)
        assert report["flag"] is True and report["gap"] == 1.0


# -------------------------------------------------------------- B2 judges


class TestJudgePanel:
    def test_order_reproducible_and_differs_across_replicas(self):
        ids = ["r1", "r2", "r3", "r4"]
        first = judges.criterion_order(ids, 20260922, "X", "openai/g", 0)
        assert first == judges.criterion_order(ids, 20260922, "X", "openai/g", 0)
        assert first != judges.criterion_order(ids, 20260922, "X", "openai/g", 1)
        assert sorted(first) == ids

    def test_heterogeneous_enforcement(self, protocol):
        judges.enforce_heterogeneous(protocol.judge.models, ["claude-sonnet-5", "claude-opus-5-5"])
        with pytest.raises(judges.JudgeError, match="stessa famiglia"):
            judges.enforce_heterogeneous(["anthropic/claude-x"], ["claude-sonnet-5"])
        with pytest.raises(judges.JudgeError, match="famiglie diverse"):
            judges.enforce_heterogeneous(["openai/a", "openai/b"], ["claude-sonnet-5"])

    def test_parse_fail_closed(self):
        assert judges.parse_residual("r1: SI\nr2: NO", ["r1", "r2"]) == {"r1": True, "r2": False}
        assert judges.parse_residual("r1: SI", ["r1", "r2"]) is None  # partial
        assert judges.parse_residual("r1: SI\nr1: NO\nr2: SI", ["r1", "r2"]) is None  # contradictory
        assert judges.parse_residual("", ["r1"]) is None

    def test_judge_item_with_stub(self):
        calls = []

        def stub(model, prompt, seed):
            calls.append((model, prompt))
            return "r1: SI\nr2: NO"

        result = judges.judge_item("X", "Q?", "A.", {"r1": "uno", "r2": "due"},
                                   ["openai/a", "google/b"], seed=1, replicas=2, call=stub)
        assert len(calls) == 4
        assert result["consensus"] == {"r1": True, "r2": False}
        assert result["intra_judge_agreement"] == 1.0
        assert result["unparseable"] == 0

    def test_failed_call_recorded_not_coerced(self):
        def broken(model, prompt, seed):
            raise judges.JudgeError("down")

        result = judges.judge_item("X", "Q?", "A.", {"r1": "uno"}, ["openai/a"], seed=1, replicas=1, call=broken)
        assert result["unparseable"] == 1 and result["consensus"] == {"r1": None}

    def test_calibration_thresholds(self, protocol):
        rows = [{"key": f"k{i}", "human": [i % 2 == 0, i % 2 == 0], "group": "A" if i < 20 else "B"} for i in range(40)]
        perfect = {f"k{i}": i % 2 == 0 for i in range(40)}
        assert judges.calibration(rows, perfect, protocol.judge.calibration)["calibrated"]
        # A judge that misses correct answers from group B only: EO gap fails.
        biased = {f"k{i}": (i % 2 == 0 and i < 20) for i in range(40)}
        report = judges.calibration(rows, biased, protocol.judge.calibration)
        assert report["eo_gap"] == pytest.approx(1.0) and not report["calibrated"]
        # Too few gold rows: never calibrated.
        assert not judges.calibration(rows[:10], perfect, protocol.judge.calibration)["calibrated"]

    def test_residual_rubrics_exclude_mechanical(self, bank):
        from benchmarks.limes.cli import _residual_rubrics

        item = bank.by_id("W1-SH-TEMPA")
        residual = _residual_rubrics(item)
        assert residual and all(not r.mechanical for r in item.rubrics if r.id in residual)
        assert len(residual) < len(item.rubrics)


# ----------------------------------------------------- markers and verdicts


class TestV1Observables:
    @pytest.mark.parametrize("text,marker,expected", [
        ("art. 25 Cost.", "cost-25", True),
        ("art. 25 c.p.", "cost-25", False),
        ("art. 25 c.p.c.", "cp-25", False),
        ("artt. 2043 e 2059 c.c.", "cc-2043", True),
        ("art. 2948, n. 4, c.c.", "cc-2948", True),
        ("Si veda l'art. 2043. Il c.c. dispone", "cc-2043", False),
        ("art. 1176 c.c. e di Corte cost.", "cost-1176", False),
        ("art. 5-bis d.lgs. 28/2010", "dlgs28-5", False),
        ("art. 1 della legge n. 689 del 1981", "l689-1", True),
        ("Cass., SS.UU., 30/12/2021, n. 41994", "cass-41994-2021", True),
        ("n. 41994/2020", "cass-41994-2021", False),
    ])
    def test_coded_markers(self, text, marker, expected):
        assert has_marker(text, marker) is expected

    def test_unknown_ids(self):
        assert not is_known_marker("xx-12")
        assert not is_known_marker("cc-")

    @pytest.mark.parametrize("text,expected", [
        ("...\n**ESITO: SÌ**", "affermare"),
        ("Esito: NO.", "negare"),
        ("ESITO: controverso", "segnalare"),
        ("ESITO: SI\nESITO: NO", "ambiguo"),
    ])
    def test_explicit_verdict(self, text, expected):
        assert detect_orientation(text) == expected

    def test_no_verdict_falls_back_to_lexicon(self):
        assert explicit_verdict("è pacifico") is None
        assert detect_orientation("è pacifico") == "affermare"


# -------------------------------------------------------- C: gate on bank


class TestCeilingAndDiscrimination:
    @pytest.fixture(scope="class")
    def sim(self):
        return simulate("wave-1")

    def test_ideal_scores_100_everywhere(self, sim):
        card = sim["models"]["ideale"]["scorecard"]
        for key in "CPHAUM":
            assert card.dimensions[key].rate == 1.0, key
        assert sim["models"]["ideale"]["discrimination"]["rate"] == 1.0

    def test_always_affirm_never_discriminates(self, sim):
        row = sim["models"]["afferma-sempre"]
        assert row["discrimination"]["rate"] == 0.0
        for key in "HAU":
            assert row["scorecard"].dimensions[key].rate < 0.6, key

    def test_engineered_failures_hit_their_dimension_only(self, sim):
        cards = {k: v["scorecard"].dimensions for k, v in sim["models"].items()}
        assert cards["calcolo-sbagliato"]["C"].rate == 0.0
        assert all(cards["calcolo-sbagliato"][k].rate == 1.0 for k in "PHAUM")
        assert cards["articolo-contiguo"]["P"].rate < 0.8
        assert cards["articolo-contiguo"]["P"].detail["disqualifying_violations"] > 0
        assert all(cards["articolo-contiguo"][k].rate == 1.0 for k in "CHAUM")
        assert all(cards["senza-fonti"][k].rate == 0.0 for k in "PHAM")
        assert cards["senza-fonti"]["C"].rate == 1.0


class TestMalformedItemsRejected:
    """Five deliberately malformed items: loader/validator refuse them with
    a message that names the item and the problem (no bare traceback)."""

    @staticmethod
    def _valid():
        return json.loads((BANK_DIR / "private" / "private-wave-1.json").read_text(encoding="utf-8"))["items"][0]

    def _bank_with(self, tmp_path, bad: dict | str) -> Path:
        root = tmp_path / "bank"
        (root / "anchors").mkdir(parents=True)
        (root / "private").mkdir()
        (root / "anchors" / "a.json").write_text(
            (BANK_DIR / "anchors" / "anchors-wave-0.json").read_text(encoding="utf-8"), encoding="utf-8")
        (root / "validity").mkdir()
        (root / "validity" / "v.json").write_text(
            (BANK_DIR / "validity" / "anchors-wave-0.json").read_text(encoding="utf-8"), encoding="utf-8")
        payload = bad if isinstance(bad, str) else json.dumps({"items": [bad]})
        (root / "private" / "p.json").write_text(payload, encoding="utf-8")
        return root

    def _expect(self, tmp_path, bad, match):
        root = self._bank_with(tmp_path, bad)
        with pytest.raises(ValidationError, match=match):
            validate_wave_slice(root, ("anchors/a.json", "private/p.json"), True)

    def test_typo_in_key(self, tmp_path):
        bad = self._valid()
        bad["validty"] = bad.pop("validity")
        self._expect(tmp_path, bad, r"unknown keys \['validty'\]")

    def test_incomplete_validity_card(self, tmp_path):
        bad = self._valid()
        bad["validity"] = {**bad["validity"], "source": ""}
        bad.pop("twin")
        bad["correct_orientation"] = "negare"
        self._expect(tmp_path, bad, "validity card incomplete")

    def test_orphan_twin(self, tmp_path):
        self._expect(tmp_path, self._valid(), "twin family")

    def test_unknown_marker(self, tmp_path):
        bad = self._valid()
        bad.pop("twin")
        bad["correct_orientation"] = "negare"
        bad["expected_markers"] = ["cc-9x"]
        self._expect(tmp_path, bad, "unknown citation marker 'cc-9x'")

    def test_construct_def_of_another_construct(self, tmp_path):
        bad = self._valid()
        bad.pop("twin")
        bad["correct_orientation"] = "negare"
        bad["validity"] = {**bad["validity"], "construct_def": "U.segnale_incertezza"}
        self._expect(tmp_path, bad, "defines 'revirement'")

    def test_invalid_json_names_the_file(self, tmp_path):
        self._expect(tmp_path, "{not json", r"p\.json: invalid JSON")

    def test_validator_cli_reports_without_traceback(self, tmp_path, capsys, monkeypatch):
        from benchmarks.limes.bank.schema import validate

        bad = self._valid()
        bad["validty"] = bad.pop("validity")
        root = self._bank_with(tmp_path, bad)
        monkeypatch.setattr("sys.argv", ["validate", str(root)])
        assert validate.main() == 1
        err = capsys.readouterr().err
        assert err.startswith("BANK INVALID:") and "Traceback" not in err


# ----------------------------------------------------------------- runner


class TestRunnerV1:
    def test_plugin_argv(self):
        argv = build_argv("plugin", "m", 30, "default", Path("/x/mcp.json"), "P",
                          plugin_dir=Path("/x/plugin"), tools=("ToolSearch", "Skill"))
        assert argv[argv.index("--plugin-dir") + 1] == "/x/plugin"
        assert argv[argv.index("--mcp-config") + 1] == "/x/mcp.json"
        assert argv[argv.index("--tools") + 1] == "ToolSearch,Skill"
        assert "--disable-slash-commands" not in argv  # the plugin's skills are the enhancement
        assert "--strict-mcp-config" in argv  # no account connectors in the cell

    def test_stream_json_and_parity(self):
        bare = build_argv("bare", "m", 30, "default", None, "P")
        assert bare[bare.index("--output-format") + 1] == "stream-json" and "--verbose" in bare
        assert bare[bare.index("--tools") + 1] == ""
        assert "--disable-slash-commands" in bare
        with pytest.raises(Exception):
            build_argv("plugin", "m", 30, "default", None, "P")  # no plugin dir: refuse

    def test_jsonl_stream_prefers_final_result(self):
        stream = "\n".join(json.dumps(m) for m in (
            {"type": "system", "subtype": "init"},
            {"type": "assistant", "message": {"content": [{"type": "text", "text": "cerco la norma"},
                                                          {"type": "tool_use", "name": "cite_law"}]}},
            {"type": "user", "message": {"content": [{"type": "tool_result", "content": "Art. 2043 testo"}]}},
            {"type": "result", "subtype": "success", "result": "Risposta: art. 2043 c.c.",
             "num_turns": 2, "total_cost_usd": 0.1},
        ))
        payload = _parse_json_output(stream)
        answer, tools, stop = _extract_payload(payload)
        assert answer == "Risposta: art. 2043 c.c."  # not the tool-call chatter
        assert tools == ["Art. 2043 testo"] and stop == "success"

    def test_freeze_guard_detects_drift(self, tmp_path):
        from benchmarks.limes.runner.shas import Shas
        from benchmarks.limes.runner.waves import Wave

        repo = tmp_path / "r"
        (repo / "bank").mkdir(parents=True)
        run = lambda *a: subprocess.run(["git", "-c", "user.email=t@t", "-c", "user.name=t", *a],
                                        cwd=repo, check=True, capture_output=True)
        run("init", "-q")
        (repo / "bank" / "x.json").write_text("[]", encoding="utf-8")
        run("add", ".")
        run("commit", "-qm", "b")
        run("tag", "limes-w")
        w = Wave(id="w", tag="limes-w", configs=["c"], models=["m"], expected_cells=0)
        shas = Shas("b", "p", "m", "c", "none")
        assert freeze_guard(w, shas, repo, [repo / "bank"])["checked"] == ["bank"]
        (repo / "bank" / "x.json").write_text("[1]", encoding="utf-8")
        with pytest.raises(WaveError, match="diverge"):
            freeze_guard(w, shas, repo, [repo / "bank"])

    def test_materialize_ref(self, tmp_path):
        from benchmarks.limes.runner.refs import materialize

        repo = tmp_path / "r"
        (repo / "plugin").mkdir(parents=True)
        run = lambda *a: subprocess.run(["git", "-c", "user.email=t@t", "-c", "user.name=t", *a],
                                        cwd=repo, check=True, capture_output=True)
        run("init", "-q")
        (repo / "plugin" / "a.txt").write_text("v1", encoding="utf-8")
        run("add", ".")
        run("commit", "-qm", "v1")
        run("tag", "v1")
        (repo / "plugin" / "a.txt").write_text("v2", encoding="utf-8")
        run("commit", "-qam", "v2")
        out = materialize("v1", repo, tmp_path / "cache")
        assert (out / "a.txt").read_text() == "v1"  # the ref, never the worktree
        assert materialize("v1", repo, tmp_path / "cache") == out  # cached


class TestRateLimits:
    """An account quota is not a property of the cell: waited out, never
    spent from the retry budget, never an exclusion."""

    @staticmethod
    def _limit(reset=None):
        text = "Claude AI usage limit reached" + (f"|{reset}" if reset else "")
        return json.dumps({"type": "result", "subtype": "success", "is_error": True, "result": text})

    @pytest.fixture()
    def env(self, monkeypatch):
        from benchmarks.limes.runner import executor
        from benchmarks.limes.runner.manifests import load_manifest

        monkeypatch.setattr(executor, "bootstrap_credentials", lambda scratch: None)
        return executor, load_manifest(LIMES / "configs" / "bare.yaml")

    def test_detection(self):
        from benchmarks.limes.runner.executor import _parse_json_output, rate_limit_reset

        assert rate_limit_reset(self._limit(), "", _parse_json_output(self._limit())) == 900.0
        ok = json.dumps({"type": "result", "subtype": "success", "result": "risposta"})
        # stderr noise never turns a real answer into a refusal
        assert rate_limit_reset(ok, "overloaded, retrying", _parse_json_output(ok)) is None

    def test_waits_and_retries_without_spending_budget(self, env, tmp_path, protocol):
        executor, manifest = env
        replies = [self._limit(), self._limit(),
                   json.dumps({"type": "result", "subtype": "success", "result": "ok"})]
        sleeps = []
        executor_run = lambda argv, **kw: subprocess.CompletedProcess(argv, 0, stdout=replies.pop(0), stderr="")
        import unittest.mock as mock

        with mock.patch.object(executor.subprocess, "run", executor_run):
            outcome = executor.run_cell(manifest=manifest, model="m", item_id="X", prompt="p",
                                        protocol=protocol, out_dir=tmp_path / "x", sleep=sleeps.append)
        assert outcome.ok and outcome.attempts == 1 and not outcome.excluded
        assert sleeps == [900.0, 900.0]

    def test_aborts_past_max_wait(self, env, tmp_path, protocol):
        executor, manifest = env
        import unittest.mock as mock

        with mock.patch.object(executor.subprocess, "run",
                               lambda argv, **kw: subprocess.CompletedProcess(argv, 0, stdout=self._limit(), stderr="")):
            with pytest.raises(executor.RateLimited, match="--resume"):
                executor.run_cell(manifest=manifest, model="m", item_id="X", prompt="p", protocol=protocol,
                                  out_dir=tmp_path / "x", sleep=lambda s: None, max_rate_limit_wait_s=1000)
        assert not (tmp_path / "x" / "config-dir").exists()

    # Shape of a REAL Pro-plan refusal (MVP run, 2026-09-23): subtype says
    # "success", the refusal text sits in `result`, the only reliable
    # signals are is_error / api_error_status / error=rate_limit. The first
    # MVP run scored this text as an answer.
    REAL_REFUSAL = "\n".join(json.dumps(m) for m in (
        {"type": "assistant", "error": "rate_limit", "is_api_error_message": True,
         "message": {"model": "<synthetic>", "role": "assistant",
                     "content": [{"type": "text", "text": "You've hit your session limit · resets 1:30pm (Europe/Rome)"}]}},
        {"type": "result", "subtype": "success", "is_error": True, "api_error_status": 429,
         "result": "You've hit your session limit · resets 1:30pm (Europe/Rome)"},
    ))

    def test_real_refusal_waits_until_announced_reset(self):
        from datetime import datetime
        from zoneinfo import ZoneInfo

        from benchmarks.limes.runner.executor import rate_limit_reset

        now = datetime(2026, 9, 23, 11, 44, tzinfo=ZoneInfo("Europe/Rome")).timestamp()
        wait = rate_limit_reset(self.REAL_REFUSAL, "", _parse_json_output(self.REAL_REFUSAL), now=now)
        assert wait == pytest.approx(107 * 60)

    def test_is_error_result_is_never_an_answer(self, env, tmp_path, protocol):
        executor, manifest = env
        import unittest.mock as mock

        other_error = json.dumps({"type": "result", "subtype": "success", "is_error": True,
                                  "result": "API Error: 500 internal"})
        with mock.patch.object(executor.subprocess, "run",
                               lambda argv, **kw: subprocess.CompletedProcess(argv, 0, stdout=other_error, stderr="")):
            outcome = executor.run_cell(manifest=manifest, model="m", item_id="X", prompt="p",
                                        protocol=protocol, out_dir=tmp_path / "x", sleep=lambda s: None)
        assert not outcome.ok and outcome.excluded and outcome.answer == ""


class TestHumanReview:
    """Review round trip: balanced export, pre-registered inclusion rule,
    twin families excluded whole, power re-checked on what survives."""

    def test_round_trip(self, tmp_path, capsys):
        from openpyxl import load_workbook

        from benchmarks.limes.bank import review

        out = tmp_path / "rev"
        assert review.main(["export", "--wave", "wave-1", "--reviewers", "A,B,C,D", "--out", str(out)]) == 0
        manifest = json.loads((out / "assegnazioni.json").read_text(encoding="utf-8"))
        loads = [len(v["items"]) for v in manifest["reviewers"].values()]
        assert max(loads) - min(loads) <= 4  # balanced
        # Every S item reviewed exactly twice; twins always by the same pair.
        seen = {}
        for reviewer, row in manifest["reviewers"].items():
            for item_id in row["items"]:
                seen.setdefault(item_id, set()).add(reviewer)
        assert all(len(r) == 2 for r in seen.values())
        assert seen["W1-SH-TEMPA"] == seen["W1-SH-TEMPB"]

        # Fill: everything OK except one ERRORE on a twin and one DA RIVEDERE.
        for reviewer, row in manifest["reviewers"].items():
            path = out / row["file"]
            wb = load_workbook(path)
            sheet = wb["Item"]
            for r in range(2, sheet.max_row + 1):
                item_id = sheet.cell(r, 1).value
                verdict = "OK"
                if item_id == "W1-SA-PENA" and reviewer == sorted(seen["W1-SA-PENA"])[0]:
                    verdict = "ERRORE"
                if item_id == "W1-SP-2043":
                    verdict = "DA RIVEDERE"
                sheet.cell(r, 10).value = verdict
            wb.save(path)
        assert review.main(["ingest", "--wave", "wave-1", "--dir", str(out)]) == 0
        decisions = json.loads((out / "decisioni.json").read_text(encoding="utf-8"))
        assert {"W1-SA-PENA", "W1-SA-PENB", "W1-SP-2043"} <= set(decisions["excluded"])
        assert "W1-SH-TEMPA" not in decisions["excluded"]
        assert "potenza OK" in capsys.readouterr().out

    def test_half_family_exclusion_refused(self, wave):
        with pytest.raises(ValidationError, match="half-excluded"):
            load_bank(BANK_DIR, wave.bank_slices, wave.require_validity, ["W1-SA-PENA"])

    def test_anchor_exclusion_refused(self, wave):
        with pytest.raises(ValidationError, match="anchors cannot be excluded"):
            load_bank(BANK_DIR, wave.bank_slices, wave.require_validity, ["QA-01", "W1-PQA-01"])
