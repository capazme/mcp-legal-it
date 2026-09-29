"""Structural live gate for the credit-recovery DOCX generators (src/tools/procure_quotazioni.py).

Benchmark phase 4, strategy "strutturale", group ``procure_quotazioni``: the benchmark site
has no calculator that produces these documents (its procura page is a free-text editor and
its fee pages are notule, not quotation letters), so the two DOCX are checked against their
own contract, against the norms they cite and against the primary source of every figure.

genera_procura_liti_docx (docstring ``Vigenza: art. 83 c.p.c.; art. 18, co. 5, D.M. 44/2011``)

* the three plan cases (docs/benchmark/piano-benchmark-andreani.json): two defenders
  ("congiuntamente e disgiuntamente", one autentica line each), one defender with fax,
  no defender (error, no file); plus a missing controparte (error) and the one-page layout
  (DOCX -> PDF with LibreOffice, pages counted with pdfinfo);
* art. 83, co. 3, c.p.c. (the autografia of the party certified by the defender) and the
  information the client declares to have received: art. 4, co. 3, D.Lgs. 28/2010,
  art. 2, co. 7, and art. 3 D.L. 132/2014, art. 13, co. 5, and art. 12, co. 1, L. 247/2012;
* art. 18, co. 5, D.M. 44/2011 "come sostituito dal D.M. 48/2013": the vigente comma (after
  D.M. 217/2023 and its avviso di rettifica, GU 15-01-2024 n. 11) is the text of D.M. 48/2013;
* the references cited in the text, extracted with a regex and passed to verifica_citazioni;
* the privacy clause: the GDPR citation form and the legacy lexicon ("dati sensibili",
  bridged by art. 22, co. 2, D.Lgs. 101/2018).

genera_quotazione_docx (docstring ``Vigenza: D.M. 55/2014 agg. D.M. 147/2022; CU ex DPR 115/2002``)

* the four plan cases: letter structure, prospetto arithmetic recomputed from the norm
  (art. 2, co. 2, D.M. 55/2014 spese generali 15%; art. 16 DPR 633/1972 IVA 22%; art. 25
  DPR 600/1973 ritenuta 20%), contributo unificato (art. 13, co. 1-3, DPR 115/2002) and the
  27 euro anticipazione forfettaria (art. 30 DPR 115/2002);
* the fee tables against the primary source: D.M. 13 agosto 2022, n. 147 as published in
  GU Serie Generale n. 236 of 8-10-2022 (codice redazionale 22G00157, allegato "Nuove tabelle
  parametri forensi", read from gazzettaufficiale.it): tabella 2 (tribunale, used for the
  opposizione), tabella 8 (procedimenti monitori), tabelle 16 and 17 (esecuzioni);
* art. 4, co. 1 e co. 1-bis, D.M. 55/2014 as vigente on Normattiva (minimum = medio -50%;
  the PCT increase is "fino al 30 per cento" since D.M. 147/2022);
* the estremi of D.M. 147/2022 cited in the letter (GU n. 236, in force from 23-10-2022) and
  the absence of a later regolamento on the parametri forensi in the GU (2023 to today).

The DOCX go to pytest's tmp_path (``_OUTPUT_DIR`` is patched): nothing is written in the
repository or in the user's temp folder. The date is pinned with ``LEGAL_TODAY``.

A failing test here is a genuine divergence of the tool from the norm (or a change of the
norm): do not soften it. It needs the network (Normattiva, EUR-Lex, gazzettaufficiale.it)
and, for the page count, LibreOffice (soffice) and pdfinfo; it is excluded from the default
suite:

    .venv/bin/pytest tests/unit/test_strutturale_procure_quotazioni_live.py -m live -q -p no:cacheprovider -rfEs
"""

from __future__ import annotations

import asyncio
import json
import re
import shutil
import subprocess
from datetime import date, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt

import src.server  # noqa: F401  (registers every tool module first: avoids a circular import)
from src.tools import procure_quotazioni as pq
from tests.unit._norme_live import assert_parole, cite_law_json, contiene, normalizza

pytestmark = pytest.mark.live


def _fn(tool):
    return getattr(tool, "fn", tool)


PROCURA = _fn(pq.genera_procura_liti_docx)
QUOTAZIONE = _fn(pq.genera_quotazione_docx)

OGGI = "2026-09-25"

# ---------------------------------------------------------------------------
# Plan cases, copied verbatim from docs/benchmark/piano-benchmark-andreani.json
# ---------------------------------------------------------------------------

PROCURA_DUE_DIFENSORI = {
    "mandante_denominazione": "Esempio S.r.l.",
    "mandante_sede": "via Roma n. 1, 20100 Milano (MI)",
    "mandante_cf_piva": "01234567890",
    "firmatario_nome": "Mario Rossi",
    "firmatario_cf": "RSSMRA70A01F205X",
    "controparte": "Delta S.r.l., in persona del legale rappresentante pro tempore, con sede in Torino, "
                   "P.IVA 09876543210",
    "difensori": [{"nome": "Giulia Bianchi", "cf": "BNCGLI80A41F205Y"}, {"nome": "Paolo Verdi", "cf": "VRDPLA75B02F205Z"}],
    "domicilio_studio": "corso Esempio n. 10, 20100 Milano (MI)",
    "pec": "g.bianchi@pec.esempio.it e p.verdi@pec.esempio.it",
    "firmatario_qualifica": "legale rappresentante",
    "fax": "",
    "luogo": "Milano",
    "data_documento": "25/09/2026",
}

PROCURA_DIFENSORE_UNICO = {
    "mandante_denominazione": "Esempio S.r.l.",
    "mandante_sede": "via Roma n. 1, 20100 Milano (MI)",
    "mandante_cf_piva": "01234567890",
    "firmatario_nome": "Mario Rossi",
    "firmatario_cf": "RSSMRA70A01F205X",
    "controparte": "Delta S.r.l., con sede in Torino, P.IVA 09876543210",
    "difensori": [{"nome": "Giulia Bianchi", "cf": "BNCGLI80A41F205Y"}],
    "domicilio_studio": "corso Esempio n. 10, 20100 Milano (MI)",
    "pec": "g.bianchi@pec.esempio.it",
    "firmatario_qualifica": "presidente del consiglio di amministrazione e legale rappresentante",
    "fax": "02 1234567",
    "luogo": "Torino",
    "data_documento": "01/08/2026",
}

PROCURA_SENZA_DIFENSORI = {
    "mandante_denominazione": "Esempio S.r.l.",
    "mandante_sede": "via Roma n. 1, 20100 Milano (MI)",
    "mandante_cf_piva": "01234567890",
    "firmatario_nome": "Mario Rossi",
    "firmatario_cf": "RSSMRA70A01F205X",
    "controparte": "Delta S.r.l.",
    "difensori": [],
    "domicilio_studio": "corso Esempio n. 10, 20100 Milano (MI)",
    "pec": "g.bianchi@pec.esempio.it",
}

_QUOTAZIONE_BASE = {
    "cliente_denominazione": "Esempio S.r.l.",
    "cliente_indirizzo": "via Roma n. 1; 20100 - Milano",
    "accettazione_denominazione": "",
    "luogo": "Milano",
    "data_documento": "25/09/2026",
    "contributo_unificato": -1,
    "compenso_fase_introduttiva": 166,
    "compenso_fase_trattazione": 284,
}

QUOT_MONITORIO_10000_MINIMI = {
    **_QUOTAZIONE_BASE, "tipo": "monitorio", "valore_causa": 10000, "debitore": "Delta S.r.l.",
    "difensori": ["Avv. Giulia Bianchi", "Avv. Paolo Verdi"], "livello": "minimi",
}
QUOT_ESECUZIONE_2500 = {
    **_QUOTAZIONE_BASE, "tipo": "esecuzione", "valore_causa": 2500, "debitore": "Delta S.r.l.",
    "difensori": ["Avv. Giulia Bianchi"], "livello": "minimi",
}
QUOT_OPPOSIZIONE_26000_MEDI = {
    **_QUOTAZIONE_BASE, "tipo": "opposizione", "valore_causa": 26000, "debitore": "Gamma S.r.l.",
    "difensori": ["Avv. Giulia Bianchi", "Avv. Paolo Verdi"], "livello": "medi",
}
QUOT_MONITORIO_1000_MEDI = {
    **_QUOTAZIONE_BASE, "tipo": "monitorio", "valore_causa": 1000, "debitore": "Delta S.r.l.",
    "difensori": ["Avv. Giulia Bianchi"], "livello": "medi",
}

#: Figures of the plan's "atteso" column (prose turned into numbers).
ATTESO_PIANO = {
    "monitorio_10000_minimi": {
        # Corrected against the norm: Tabella 8 medio 567, minimo = 567 x 50% = 283,50 exact
        # (art. 4, co. 1, D.M. 55/2014). The plan's 284,00 was the euro-rounded minimum.
        "Compenso tabellare": "283,50", "Totale variazioni in aumento": "85,05", "Compenso totale": "368,55",
        "Spese generali": "55,28", "Cassa Avvocati": "16,95", "Totale imponibile": "440,78",
        "IVA 22%": "96,97", "IPOTESI DI COMPENSO LIQUIDABILE": "537,75", "A dedurre ritenuta": "84,77",
        "Totale documento": "452,98",
    },
    "esecuzione_2500": {
        # Tabella 17, 1.100,01-5.200: minimi 165,50 + 283,50 = 449,00 (the plan's 450 was rounded)
        "Compenso tabellare": "449,00", "Spese generali": "67,35", "Cassa Avvocati": "20,65",
        "Totale imponibile": "537,00", "IVA 22%": "118,14", "IPOTESI DI COMPENSO LIQUIDABILE": "655,14",
    },
    "opposizione_26000_medi": {
        "Compenso tabellare": "5.077,00", "Totale variazioni in aumento": "1.523,10",
        "Spese generali": "990,02", "Cassa Avvocati": "303,60", "Totale imponibile": "7.893,72",
        "IVA 22%": "1.736,62", "IPOTESI DI COMPENSO LIQUIDABILE": "9.630,34",
        "A dedurre ritenuta": "1.518,02", "Totale documento": "8.112,32",
    },
    "monitorio_1000_medi": {"Compenso tabellare": "473,00"},
}

# D.M. 13 agosto 2022, n. 147 in the Gazzetta Ufficiale (Serie Generale).
GU_DM_147_CODICE = "22G00157"
GU_DM_147_DATA = "2022-10-08"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


@pytest.fixture
def out_dir(tmp_path, monkeypatch) -> Path:
    """Redirect the tool's output folder to tmp_path and pin today's date."""
    cartella = tmp_path / "docx"
    monkeypatch.setattr(pq, "_OUTPUT_DIR", str(cartella))
    monkeypatch.setenv("LEGAL_TODAY", OGGI)
    return cartella


def _euro(testo: str) -> Decimal:
    """'€ 1.234,56' (or '+ € 85,20') -> Decimal('1234.56')."""
    m = re.search(r"(\d{1,3}(?:\.\d{3})*,\d{2})", testo)
    assert m, f"nessun importo in {testo!r}"
    return Decimal(m.group(1).replace(".", "").replace(",", "."))


def _d2(valore) -> Decimal:
    return Decimal(str(valore)).quantize(Decimal("0.01"), ROUND_HALF_UP)


def _apri(risultato: str) -> tuple[Path, Document]:
    m = re.search(r"File salvato: (.+?\.docx) \(", risultato)
    assert m, risultato
    percorso = Path(m.group(1))
    assert percorso.is_file(), percorso
    return percorso, Document(str(percorso))


def _testo(doc: Document) -> str:
    parti = [p.text for p in doc.paragraphs]
    for tabella in doc.tables:
        for riga in tabella.rows:
            parti.append(" | ".join(cella.text for cella in riga.cells))
    return "\n".join(parti)


def _prospetto(doc: Document) -> dict[str, str]:
    """First table of the quotation: {voce: importo as printed}."""
    righe = {}
    for riga in doc.tables[0].rows:
        voce, importo = (c.text.strip() for c in riga.cells)
        righe[voce] = importo
    return righe


def _voce(prospetto: dict[str, str], prefisso: str) -> Decimal:
    """Amount of the row named `prefisso` (exact name first, else the only row starting with it)."""
    if prefisso in prospetto:
        return _euro(prospetto[prefisso])
    trovate = [v for k, v in prospetto.items() if k.startswith(prefisso)]
    assert len(trovate) == 1, (prefisso, prospetto)
    return _euro(trovate[0])


def _catena(tabellare: Decimal, aumento_pct: bool, ritenuta: bool = True) -> dict[str, Decimal]:
    """Prospetto recomputed from the norm, rounding each line to the cent.

    art. 4, co. 1-bis, D.M. 55/2014 (increase for PCT drafting, up to 30%: the tool applies
    the maximum); art. 2, co. 2, D.M. 55/2014 (spese forfettarie 15% del compenso totale);
    contributo integrativo Cassa Forense 4% on compenso + spese generali; art. 16 DPR
    633/1972 (IVA 22%); art. 25 DPR 600/1973 (ritenuta 20% on compenso + spese generali).
    """
    aumento = _d2(tabellare * Decimal("0.30")) if aumento_pct else Decimal("0.00")
    totale = _d2(tabellare + aumento)
    spese_generali = _d2(totale * Decimal("0.15"))
    cpa = _d2((totale + spese_generali) * Decimal("0.04"))
    imponibile = totale + spese_generali + cpa
    iva = _d2(imponibile * Decimal("0.22"))
    liquidabile = imponibile + iva
    out = {
        "Spese generali": spese_generali, "Cassa Avvocati": cpa, "Totale imponibile": imponibile,
        "IVA 22%": iva, "IPOTESI DI COMPENSO LIQUIDABILE": liquidabile,
    }
    if aumento_pct:
        out["Totale variazioni in aumento"] = aumento
        out["Compenso totale"] = totale
    if ritenuta:
        out["A dedurre ritenuta"] = _d2((totale + spese_generali) * Decimal("0.20"))
        out["Totale documento"] = liquidabile - out["A dedurre ritenuta"]
    return out


def _verifica_citazioni(citazioni: list[str]) -> list[dict]:
    from src.tools.legal_citations import verifica_citazioni

    raw = asyncio.run(_fn(verifica_citazioni)(citazioni="\n".join(citazioni), formato="json"))
    payload = json.loads(raw)
    assert payload.get("errore") is None, payload
    return payload["citazioni"]


# --- normative references cited in the procura --------------------------------------------

_ATTO = (
    r"(c\.p\.c\."
    r"|D\.Lgs\.\s*n\.\s*\d+/\d{4}"
    r"|D\.L\.\s*n\.\s*\d+/\d{4}"
    r"|D\.M\.(?:\s*Giustizia)?\s*n\.\s*\d+/\d{4})"
)
_ART_ESPLICITO = re.compile(
    r"\bart\.\s*(\d+)((?:,\s*(?:co\.|comma)\s*\d+(?:[-\s]?bis)?)?),?\s*(?:del\s+)?" + _ATTO, re.I
)
_ART_RELATIVO = re.compile(
    r"\bartt?\.\s*(\d+(?:\s*(?:e|,)\s*(?:\d+|ss\.))*)\s+del\s+(?:medesimo|suddetto)\s+decreto", re.I
)
_ATTO_CITATO = re.compile(_ATTO, re.I)


def _normalizza_atto(atto: str) -> str:
    atto = re.sub(r"\s+", " ", atto)
    if atto.lower() == "c.p.c.":
        return "c.p.c."
    # verifica_citazioni does not parse "D.M. Giustizia n. 44/2011": the issuing ministry
    # adds nothing to the identification of the act, so it is dropped.
    return atto.replace("D.M. Giustizia", "D.M.")


def _riferimenti_procura(testo: str) -> list[str]:
    """Article-level references of the procura, relative ones resolved to the act cited before."""
    trovati: list[tuple[int, str]] = []
    for m in _ART_ESPLICITO.finditer(testo):
        comma = re.sub(r"\s+", " ", m.group(2)).lower()
        trovati.append((m.start(), f"art. {m.group(1)}{comma}, {_normalizza_atto(m.group(3))}"))
    for m in _ART_RELATIVO.finditer(testo):
        precedenti = [a for a in _ATTO_CITATO.finditer(testo[: m.start()])]
        assert precedenti, m.group(0)
        atto = _normalizza_atto(precedenti[-1].group(1))
        for numero in re.findall(r"\d+", m.group(1)):
            trovati.append((m.start(), f"art. {numero} {atto}"))
    visti, out = set(), []
    for _, rif in sorted(trovati):
        if rif not in visti:
            visti.add(rif)
            out.append(rif)
    return out


# --- tables of the D.M. 147/2022 allegato in the Gazzetta Ufficiale -------------------------

_IMPORTO = re.compile(r"\d{1,3}(?:[.,]\d{3})*,\d{2}")


def _importo_gu(testo: str) -> Decimal:
    # The GU prints thousands with a dot, and once with a comma ("52,000,00").
    intero, cent = testo[:-3], testo[-2:]
    return Decimal(re.sub(r"[.,]", "", intero) + "." + cent)


def _tabella_gu(testo: str, numero: int) -> tuple[list[Decimal], dict[str, list[Decimal]]]:
    """Parse table n. `numero` of the allegato: (upper bounds of the scaglioni, {fase: medi})."""
    allegato = testo[testo.index("Nuove tabelle parametri forensi"):]
    inizio = re.search(rf"^{numero}\. [A-Z]{{3,}}", allegato, re.M)
    assert inizio, f"tabella {numero} non trovata nell'allegato GU"
    resto = allegato[inizio.end():]
    fine = re.search(rf"^{numero + 1}\. [A-Z]{{3,}}", resto, re.M)
    sezione = resto[: fine.start()] if fine else resto
    blocchi, corrente = [], []
    for riga in sezione.splitlines():
        if riga.startswith("+"):
            if corrente:
                blocchi.append(corrente)
                corrente = []
        elif riga.startswith("|"):
            corrente.append(riga)
    if corrente:
        blocchi.append(corrente)

    def colonne(blocco: list[str]) -> list[str]:
        cols: list[str] | None = None
        for riga in blocco:
            celle = riga.strip().strip("|").split("|")
            if cols is None:
                cols = [""] * len(celle)
            for i, cella in enumerate(celle[: len(cols)]):
                cols[i] += " " + cella.strip()
        return [re.sub(r"-\s+", "", re.sub(r"\s+", " ", c)).strip() for c in cols or []]

    intestazione = colonne(blocchi[0])
    soglie = [_importo_gu(_IMPORTO.findall(c)[-1]) for c in intestazione[1:]]
    fasi = {}
    for blocco in blocchi[1:]:
        cols = colonne(blocco)
        fasi[cols[0]] = [_importo_gu(_IMPORTO.findall(c)[0]) for c in cols[1:]]
    return soglie, fasi


def _fase(fasi: dict[str, list[Decimal]], chiave: str) -> list[Decimal]:
    trovate = [v for k, v in fasi.items() if chiave in k.lower()]
    assert len(trovate) == 1, (chiave, list(fasi))
    return trovate[0]


@pytest.fixture(scope="module")
def dm147_gu() -> str:
    """Full text of D.M. 147/2022 with its allegato, as published in GU n. 236 of 8-10-2022."""
    from src.lib.gazzetta import client as gu
    from src.tools.gazzetta import _resolve_serie_path

    atto = asyncio.run(gu.fetch_atto(_resolve_serie_path("serie_generale"), GU_DM_147_DATA, GU_DM_147_CODICE))
    assert "Nuove tabelle parametri forensi" in atto.text, atto.text[:500]
    return atto.text


# ===========================================================================
# genera_procura_liti_docx
# ===========================================================================


def test_procura_due_difensori_struttura_piano(out_dir):
    """Plan case 1: two defenders, joint and several mandate, one autentica line per defender."""
    percorso, doc = _apri(PROCURA(**PROCURA_DUE_DIFENSORI))
    assert percorso.parent == out_dir
    testo = _testo(doc)
    paragrafi = [p.text for p in doc.paragraphs]

    mancanti = contiene(
        testo,
        "PROCURA ALLE LITI",
        "RILASCIATA AI SENSI DELL’ART. 83, COMMA 3, C.P.C.",
        "Il sottoscritto Mario Rossi (Cod. Fisc. RSSMRA70A01F205X), nella sua qualità di legale "
        "rappresentante di Esempio S.r.l., con sede legale in via Roma n. 1, 20100 Milano (MI), "
        "Cod. Fisc. e Partita IVA 01234567890",
        "CONFERISCE PROCURA AD LITEM",
        "all’avv. Giulia Bianchi (Cod. Fisc. BNCGLI80A41F205Y) e all’avv. Paolo Verdi (Cod. Fisc. VRDPLA75B02F205Z)",
        "congiuntamente e disgiuntamente tra loro",
        # the controparte clause is reproduced verbatim
        "nei confronti di Delta S.r.l., in persona del legale rappresentante pro tempore, con sede in "
        "Torino, P.IVA 09876543210",
        "in ogni successiva fase e grado",
        "Elegge domicilio presso i nominati difensori con studio in corso Esempio n. 10, 20100 Milano (MI), "
        "PEC g.bianchi@pec.esempio.it e p.verdi@pec.esempio.it.",
        "DICHIARA",
        "art. 4, co. 3, D.Lgs. n. 28/2010",
        "artt. 17 e 20",
        "art. 2, co. 7, D.L. n. 132/2014, convertito in L. n. 162/2014",
        "all’art. 3 del suddetto decreto",
        "D.Lgs. n. 196/2003",
        "art. 18, co. 5, D.M. Giustizia n. 44/2011",
        "Milano, 25 settembre 2026",
        "È vera e autentica",
    )
    assert not mancanti, mancanti
    assert "fax" not in testo.lower()

    # signature block of the mandante, then the autentica with one line per defender
    i_firma = paragrafi.index("Firma")
    assert paragrafi[i_firma + 1].startswith("Mario Rossi\t")
    i_autentica = paragrafi.index("È vera e autentica")
    righe_autentica = [p for p in paragrafi[i_autentica + 1:] if p.startswith("Avv. ")]
    assert [r.split("\t")[0] for r in righe_autentica] == ["Avv. Giulia Bianchi", "Avv. Paolo Verdi"]
    assert all("____" in r for r in righe_autentica)

    # formatting declared in the docstring: Times New Roman 11
    normale = doc.styles["Normal"].font
    assert normale.name == "Times New Roman" and normale.size == Pt(11)


def test_procura_difensore_unico_con_fax_piano(out_dir):
    """Plan case 2: one defender, fax in the recapiti, qualifica in full, date in words."""
    _, doc = _apri(PROCURA(**PROCURA_DIFENSORE_UNICO))
    testo = _testo(doc)
    paragrafi = [p.text for p in doc.paragraphs]

    assert "congiuntamente e disgiuntamente" not in testo
    mancanti = contiene(
        testo,
        "nella sua qualità di presidente del consiglio di amministrazione e legale rappresentante di Esempio S.r.l.",
        "all’avv. Giulia Bianchi (Cod. Fisc. BNCGLI80A41F205Y), per rappresentare e difendere",
        "Elegge domicilio presso il nominato difensore con studio in corso Esempio n. 10, 20100 Milano (MI), "
        "fax 02 1234567, PEC g.bianchi@pec.esempio.it.",
        "Torino, 1 agosto 2026",
        "art. 4, co. 3, D.Lgs. n. 28/2010",
        "art. 2, co. 7, D.L. n. 132/2014",
        "art. 18, co. 5, D.M. Giustizia n. 44/2011",
    )
    assert not mancanti, mancanti
    i_autentica = paragrafi.index("È vera e autentica")
    righe_autentica = [p for p in paragrafi[i_autentica + 1:] if p.startswith("Avv. ")]
    assert [r.split("\t")[0] for r in righe_autentica] == ["Avv. Giulia Bianchi"]


def test_procura_errori_senza_difensori_e_senza_controparte(out_dir):
    """Plan case 3 (no defender) and the missing controparte: an error and no file written."""
    risultato = PROCURA(**PROCURA_SENZA_DIFENSORI)
    assert risultato.startswith("Errore:") and "almeno un difensore" in risultato
    risultato = PROCURA(**{**PROCURA_DUE_DIFENSORI, "controparte": "  "})
    assert risultato.startswith("Errore:") and "controparte" in risultato
    assert not out_dir.exists() or not any(out_dir.iterdir())


def test_procura_una_sola_pagina(out_dir, tmp_path):
    """The docstring promises a one-page procura: DOCX -> PDF with LibreOffice, pages via pdfinfo."""
    soffice, pdfinfo = shutil.which("soffice"), shutil.which("pdfinfo")
    if not soffice or not pdfinfo:
        pytest.skip("LibreOffice (soffice) o pdfinfo non disponibili: conteggio pagine non eseguibile")
    pdf_dir = tmp_path / "pdf"
    for caso in (PROCURA_DUE_DIFENSORI, PROCURA_DIFENSORE_UNICO):
        percorso, _ = _apri(PROCURA(**caso))
        subprocess.run(
            [soffice, f"-env:UserInstallation=file://{tmp_path / 'lo_profile'}", "--headless",
             "--convert-to", "pdf", "--outdir", str(pdf_dir), str(percorso)],
            check=True, capture_output=True, timeout=180,
        )
        pdf = pdf_dir / (percorso.stem + ".pdf")
        info = subprocess.run([pdfinfo, str(pdf)], check=True, capture_output=True, text=True).stdout
        pagine = int(re.search(r"^Pages:\s+(\d+)", info, re.M).group(1))
        assert pagine == 1, f"{percorso.name}: {pagine} pagine"


def test_procura_art_83_co_3_cpc_autentica_del_difensore(out_dir):
    """art. 83, co. 3, c.p.c.: procura apposta in calce, autografia certified by the defender.

    The procura carries the defender's "È vera e autentica" after the party's signature and,
    for art. 83, co. 4, the express will to extend it beyond the grade ("ogni successiva fase
    e grado").
    """
    assert_parole(
        "art. 83 c.p.c.",
        "la procura speciale puo' essere anche apposta in calce|la procura speciale può essere anche apposta in calce",
        "l'autografia della sottoscrizione della parte deve essere certificata dal difensore",
        "si presume conferita soltanto per un determinato grado del processo, quando nell'atto non e' "
        "espressa volonta' diversa|si presume conferita soltanto per un determinato grado del processo, "
        "quando nell'atto non è espressa volontà diversa",
    )
    testo = _testo(_apri(PROCURA(**PROCURA_DUE_DIFENSORI))[1])
    assert "È vera e autentica" in testo and "in ogni successiva fase e grado" in testo


def test_procura_informative_al_cliente_nel_testo_vigente(out_dir):
    """The declarations match the information duties in force.

    art. 4, co. 3, D.Lgs. 28/2010 (mediation, "articoli 17 e 20", in writing, signed by the
    client); art. 2, co. 7, D.L. 132/2014 (negoziazione assistita) and art. 3 (condizione di
    procedibilita'); art. 13, co. 5, L. 247/2012 (complexity, oneri, written estimate split into
    oneri, spese and compenso); art. 12, co. 1, L. 247/2012 (estremi della polizza).
    """
    assert_parole(
        "art. 4 D.Lgs. 28/2010",
        "all'atto del conferimento dell'incarico, l'avvocato e' tenuto a informare l'assistito della "
        "possibilita' di avvalersi del procedimento di mediazione",
        "agevolazioni fiscali di cui agli articoli 17 e 20",
        "condizione di procedibilita' della domanda giudiziale",
        "l'informazione deve essere fornita chiaramente e per iscritto",
        "e' sottoscritto dall'assistito",
    )
    assert_parole(
        "art. 2 D.L. 132/2014",
        "e' dovere deontologico degli avvocati informare il cliente all'atto del conferimento "
        "dell'incarico della possibilita' di ricorrere alla convenzione di negoziazione assistita",
    )
    assert_parole("art. 3 D.L. 132/2014", "condizione di procedibilita' della domanda giudiziale")
    assert_parole(
        "art. 13 L. 247/2012",
        "rendere noto al cliente il livello della complessita' dell'incarico",
        "oneri ipotizzabili dal momento del conferimento alla conclusione dell'incarico",
        "la prevedibile misura del costo della prestazione, distinguendo fra oneri, spese, anche "
        "forfetarie, e compenso professionale",
    )
    assert_parole("art. 12 L. 247/2012", "l'avvocato rende noti al cliente gli estremi della propria polizza assicurativa")

    testo = _testo(_apri(PROCURA(**PROCURA_DUE_DIFENSORI))[1])
    mancanti = contiene(
        testo,
        "della possibilità di ricorrere al procedimento di mediazione",
        "dei benefici fiscali di cui agli artt. 17 e 20 del medesimo decreto",
        "dei casi in cui l’esperimento del procedimento di mediazione è condizione di procedibilità",
        "della possibilità di ricorrere alla convenzione di negoziazione assistita",
        "il grado di complessità dell’incarico",
        "oneri ipotizzabili dal momento del conferimento sino alla conclusione dell’incarico",
        "preventivo scritto relativo alla prevedibile misura dei costi della prestazione",
        "degli estremi della polizza assicurativa professionale",
    )
    assert not mancanti, mancanti


def test_procura_riferimenti_normativi_verifica_citazioni(out_dir):
    """Every article-level reference in the procura exists on Normattiva (verifica_citazioni)."""
    testo = _testo(_apri(PROCURA(**PROCURA_DUE_DIFENSORI))[1])
    riferimenti = _riferimenti_procura(testo)
    assert set(riferimenti) == {
        "art. 83, comma 3, c.p.c.",
        "art. 4, co. 3, D.Lgs. n. 28/2010",
        "art. 17 D.Lgs. n. 28/2010",
        "art. 20 D.Lgs. n. 28/2010",
        "art. 2, co. 7, D.L. n. 132/2014",
        "art. 2 D.L. n. 132/2014",
        "art. 3 D.L. n. 132/2014",
        "art. 18, co. 5, D.M. n. 44/2011",
    }, riferimenti
    esiti = _verifica_citazioni(riferimenti)
    non_verificate = [(e["citazione"], e["verdetto"], e.get("nota")) for e in esiti if e["verdetto"] != "verificata"]
    assert not non_verificate, non_verificate


def test_procura_art_18_co_5_dm_44_2011_come_sostituito_dal_dm_48_2013(out_dir):
    """art. 18, co. 5, D.M. 44/2011 "come sostituito dal D.M. 48/2013" is still the vigente text.

    D.M. 48/2013, art. 1, replaced art. 18; D.M. 217/2023 abrogated commi 1-3 and, after the
    avviso di rettifica (GU 15-01-2024, n. 11), no longer abrogates commi 4-6. The comma the
    procura invokes is therefore in force with the wording of D.M. 48/2013.
    """
    vigente = cite_law_json("art. 18 D.M. 44/2011")
    assert not vigente.get("errore"), vigente
    dm_48 = cite_law_json("art. 1 D.M. n. 48/2013")
    assert not dm_48.get("errore"), dm_48
    assert contiene(dm_48["testo"], "l'articolo 18 del decreto del ministro della giustizia 21 febbraio 2011 n. 44 è sostituito dal seguente") == []

    def comma_5(testo: str) -> str:
        m = re.search(r"5\. (La procura alle liti si considera apposta in calce.+?anche per immagine\.)", testo, re.S)
        assert m, testo
        return normalizza(m.group(1)).replace("è", "e'").replace("é", "e'")

    assert comma_5(vigente["testo"]) == comma_5(dm_48["testo"])
    assert contiene(vigente["testo"], "comma abrogato dal decreto 29 dicembre 2023, n. 217") == []
    assert contiene(vigente["testo"], "non prevede piu' l'abrogazione dei commi 4, 5 e 6") == []

    testo = _testo(_apri(PROCURA(**PROCURA_DUE_DIFENSORI))[1])
    assert "art. 18, co. 5, D.M. Giustizia n. 44/2011, come sostituito dal D.M. Giustizia n. 48/2013" in testo


def test_procura_privacy_regolamento_ue_citato_nella_forma_ufficiale(out_dir):
    """The privacy clause cites the GDPR as "Regolamento UE n. 679/2016".

    Since 1 January 2015 EU acts are numbered "(UE) anno/numero": the act is "Regolamento (UE)
    2016/679", the form the Italian legislator uses (art. 22 D.Lgs. 101/2018) and the one the
    plan requires. The inverted form still resolves (verifica_citazioni accepts it), so this is a
    formal defect of the citation, not a missing act.
    """
    assert_parole("art. 22 D.Lgs. 101/2018", "regolamento (ue) 2016/679")
    esiti = _verifica_citazioni(["art. 9 Regolamento (UE) 2016/679", "art. 2-sexies D.Lgs. n. 196/2003"])
    assert [e["verdetto"] for e in esiti] == ["verificata", "verificata"], esiti

    testo = _testo(_apri(PROCURA(**PROCURA_DUE_DIFENSORI))[1])
    assert "Regolamento (UE) 2016/679" in testo, (
        "la procura cita il GDPR come «Regolamento UE n. 679/2016»: la forma corretta è "
        "«Regolamento (UE) 2016/679»"
    )


def test_procura_privacy_lessico_dati_sensibili_ponte_art_22_dlgs_101_2018(out_dir):
    """"Dati sensibili" is pre-GDPR lexicon, still readable through art. 22, co. 2, D.Lgs. 101/2018.

    The comma maps "dati sensibili" onto the categorie particolari of art. 9 GDPR, so the clause
    is not void; but art. 9(2)(f) GDPR (defence of legal claims) makes the lawyer's processing
    lawful without the client's "autorizzazione", which the clause still asks for. Kept as a
    documented finding (no failing assertion): the wording is outdated, not unlawful.
    """
    assert_parole(
        "art. 22 D.Lgs. 101/2018",
        "le espressioni «dati sensibili» e «dati giudiziari»|le espressioni \"dati sensibili\" e \"dati giudiziari\"",
        "si intendono riferite, rispettivamente, alle categorie particolari di dati di cui all'articolo 9",
    )
    art_9 = cite_law_json("art. 9 Regolamento (UE) 2016/679")
    assert not art_9.get("errore"), art_9
    assert contiene(art_9["testo"], "categorie particolari di dati personali", "sede giudiziaria") == []

    testo = _testo(_apri(PROCURA(**PROCURA_DUE_DIFENSORI))[1])
    assert "i dati personali, anche sensibili" in testo
    assert "autorizzando sin d’ora il rispettivo trattamento" in testo


# ===========================================================================
# genera_quotazione_docx
# ===========================================================================


@pytest.mark.parametrize(
    "chiave, caso",
    [
        ("monitorio_10000_minimi", QUOT_MONITORIO_10000_MINIMI),
        ("esecuzione_2500", QUOT_ESECUZIONE_2500),
        ("opposizione_26000_medi", QUOT_OPPOSIZIONE_26000_MEDI),
        ("monitorio_1000_medi", QUOT_MONITORIO_1000_MEDI),
    ],
)
def test_quotazione_prospetto_casi_piano(out_dir, chiave, caso):
    """Plan cases: the prospetto equals the plan figures and the chain recomputed from the norm."""
    risultato = QUOTAZIONE(**caso)
    _, doc = _apri(risultato)
    prospetto = _prospetto(doc)
    for voce, atteso in ATTESO_PIANO[chiave].items():
        assert _voce(prospetto, voce) == _euro(atteso), (voce, prospetto)

    tabellare = _voce(prospetto, "Compenso tabellare")
    attesi = _catena(tabellare, aumento_pct=caso["tipo"] != "esecuzione", ritenuta=caso["tipo"] != "esecuzione")
    for voce, atteso in attesi.items():
        assert _voce(prospetto, voce) == atteso, (voce, prospetto)
    if caso["tipo"] == "esecuzione":
        # model choice, documented in the plan: no PCT increase and no ritenuta in the esecuzione
        assert not any(k.startswith(("Aumento", "A dedurre", "Totale documento")) for k in prospetto)


def test_quotazione_struttura_lettera(out_dir):
    """Letter structure for the three types: addressee, date, oggetto, DM, oneri, signatures, acceptance."""
    for caso, oggetto, specifico in (
        (QUOT_MONITORIO_10000_MINIMI, "Oggetto: Quotazione giudiziaria procedimento monitorio Delta S.r.l.",
         ["€ 118,50 a titolo di contributo unificato ed € 27,00 per la marca da bollo, per un totale "
          "complessivo preventivato di € 683,25", "misura fissa di € 200,00", "misura proporzionale del 3%"]),
        (QUOT_ESECUZIONE_2500, "Oggetto: Quotazione giudiziaria esecuzione forzata Delta S.r.l.",
         ["€ 139,00 a titolo di contributo unificato, € 27,00 per la marca da bollo ed € 120,00 a titolo "
          "forfettario", "totale complessivo preventivato di € 941,14", "atto di pignoramento"]),
        (QUOT_OPPOSIZIONE_26000_MEDI,
         "Oggetto: Quotazione giudiziaria giudizio di opposizione a decreto ingiuntivo Gamma S.r.l.",
         ["(art. 645 c.p.c.) instaurato da Gamma S.r.l.",
          "il contributo unificato, che nel giudizio di opposizione è a carico della parte opponente"]),
    ):
        _, doc = _apri(QUOTAZIONE(**caso))
        testo = _testo(doc)
        paragrafi = [p.text for p in doc.paragraphs]
        mancanti = contiene(
            testo,
            "Spett.le Società",
            "Esempio S.r.l.",
            "via Roma n. 1",
            "20100 - Milano",
            "Milano, 25 settembre 2026",
            oggetto,
            "D.M. 13 agosto 2022, n. 147, pubblicato nella Gazzetta Ufficiale n. 236 dell'8 ottobre 2022 "
            "ed in vigore dal 23 ottobre 2022",
            f"Valore della causa: € {caso['valore_causa']:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."),
            "Per integrale accettazione della presente quotazione",
            *specifico,
        )
        assert not mancanti, (caso["tipo"], mancanti)
        firme = [c.text for c in doc.tables[1].rows[0].cells if c.text]
        assert firme == caso["difensori"]
        i_acc = next(i for i, p in enumerate(paragrafi) if p.startswith("Per integrale accettazione"))
        assert paragrafi[i_acc + 1] == "Esempio S.r.l." and paragrafi[i_acc + 2].startswith("____")


def test_quotazione_contributo_unificato_e_anticipazione_forfettaria(out_dir):
    """CU: art. 13, co. 1, 2 and 3, DPR 115/2002; the 27 euro: art. 30 DPR 115/2002.

    Monitorio: half of co. 1 (co. 3), 237/2 = 118,50 for 10.000 and 43/2 = 21,50 for 1.000.
    Esecuzione: 278/2 = 139 for the esecuzioni other than immobiliare, 43 for the mobiliari
    "di valore inferiore a 2.500 euro" (2.500 itself pays 139, 2.499,99 pays 43).
    art. 30: the 27 euro are anticipated by "la parte che per prima si costituisce in giudizio"
    or deposits the ricorso or, in the esecuzione, asks for assignment or sale: the creditor in
    the monitorio and in the esecuzione; in the opposizione it is the opponente (see note in
    the benchmark result on the letter's "oltre € 27,00").
    """
    assert_parole(
        "art. 13 DPR 115/2002",
        "a) euro 43 per i processi di valore fino a 1.100 euro",
        "c) euro 237 per i processi di valore superiore a euro 5.200 e fino a euro 26.000",
        "per i processi di esecuzione immobiliare il contributo dovuto e' pari a euro 278. per gli altri "
        "processi esecutivi lo stesso importo e' ridotto della meta'",
        "per i processi esecutivi mobiliari di valore inferiore a 2.500 euro il contributo dovuto e' pari a euro 43",
        "il contributo e' ridotto alla meta' per i processi speciali previsti nel libro iv, titolo i, del "
        "codice di procedura civile, compreso il giudizio di opposizione a decreto ingiuntivo",
    )
    assert_parole(
        "art. 30 DPR 115/2002",
        "la parte che per prima si costituisce in giudizio, che deposita il ricorso introduttivo",
        "nella misura di euro 27",
    )

    def oneri(caso) -> str:
        return _testo(_apri(QUOTAZIONE(**caso))[1])

    assert "€ 118,50 a titolo di contributo unificato" in oneri(QUOT_MONITORIO_10000_MINIMI)
    assert "€ 21,50 a titolo di contributo unificato" in oneri(QUOT_MONITORIO_1000_MEDI)
    assert "€ 139,00 a titolo di contributo unificato" in oneri(QUOT_ESECUZIONE_2500)
    assert "€ 43,00 a titolo di contributo unificato" in oneri({**QUOT_ESECUZIONE_2500, "valore_causa": 2499.99})
    for caso in (QUOT_MONITORIO_10000_MINIMI, QUOT_ESECUZIONE_2500):
        assert "€ 27,00 per la marca da bollo" in oneri(caso)


def test_quotazione_aliquote_catena_nel_testo_vigente():
    """Rates of the chain: 15% (art. 2, co. 2, D.M. 55/2014), 22% (art. 16 DPR 633/1972), 20%
    (art. 25 DPR 600/1973, 19% -> 20% by art. 21, co. 11, L. 449/1997: Normattiva keeps the
    original "15 per cento" in the body and records the raises in the update notes)."""
    assert_parole(
        "art. 2 D.M. 10 marzo 2014, n. 55",
        "in ogni caso ed anche in caso di determinazione contrattuale",
        "nella misura del 15 per cento del compenso totale per la prestazione",
    )
    assert_parole("art. 16 DPR 633/1972", "nella misura del ventidue per cento della base imponibile")
    assert_parole(
        "art. 25 DPR 600/1973",
        "e' elevata al 19 per cento",
        'al primo comma le parole: "19 per cento" sono sostituite dalle seguenti: "20 per cento"',
    )


def test_quotazione_dm_147_2022_estremi_gu_ed_entrata_in_vigore():
    """"D.M. 13 agosto 2022, n. 147, pubblicato nella GU n. 236 dell'8 ottobre 2022 ed in vigore dal
    23 ottobre 2022": estremi from the GU search, art. 7 (15th day after publication)."""
    from src.tools.gazzetta import cerca_gazzetta_ufficiale

    esito = asyncio.run(_fn(cerca_gazzetta_ufficiale)(titolo="professione forense", anno_da="2022", anno_a="2022"))
    assert "DECRETO 13 agosto 2022, n. 147" in esito, esito
    assert "(GU n.236 del 8-10-2022)" in esito, esito
    assert GU_DM_147_CODICE in esito

    assert_parole(
        "art. 7 D.M. 13 agosto 2022, n. 147",
        "entra in vigore il quindicesimo giorno successivo a quello della sua pubblicazione",
    )
    assert date(2022, 10, 8) + timedelta(days=15) == date(2022, 10, 23)
    esiti = _verifica_citazioni(["art. 7 D.M. 13 agosto 2022, n. 147", "art. 645 c.p.c.", "art. 4 D.M. 10 marzo 2014, n. 55"])
    assert [e["verdetto"] for e in esiti] == ["verificata"] * 3, esiti


def test_nessun_regolamento_parametri_forensi_successivo_al_dm_147_2022():
    """Vintage guard: no later regolamento on the parametri forensi in the GU (2023 to today).

    Positive control first (the 2022 search finds D.M. 147/2022), then the same title search on
    the following years. A hit means the tables in the tool and in src/data need a new check.
    """
    from src.tools.gazzetta import cerca_gazzetta_ufficiale

    cerca = _fn(cerca_gazzetta_ufficiale)
    controllo = asyncio.run(cerca(titolo="professione forense", anno_da="2022", anno_a="2022"))
    assert "n. 147" in controllo, controllo
    esito = asyncio.run(cerca(titolo="professione forense", anno_da="2023", anno_a=OGGI[:4], max_risultati=50))
    assert "10 marzo 2014, n. 55" not in esito, esito


def test_quotazione_art_4_co_1bis_aumento_pct_fino_al_30_per_cento(out_dir):
    """art. 4, co. 1-bis, D.M. 55/2014: since D.M. 147/2022 the increase is "fino al 30 per cento".

    D.M. 147/2022, art. 2, co. 1, lett. b), replaced "e' di regola ulteriormente aumentato del
    30 per cento" with "e' ulteriormente aumentato fino al 30 per cento". The tool applies the
    maximum (allowed) but labels the row "Aumento del 30% ... (art. 4, co. 1 bis)", i.e. the
    pre-2022 rule; the row should say that 30% is the maximum of the range.
    """
    assert_parole(
        "art. 4 D.M. 10 marzo 2014, n. 55",
        "e' ulteriormente aumentato fino al 30 per cento quando gli atti depositati con modalita' telematiche",
        "possono essere diminuiti in ogni caso non oltre il 50 per cento",
    )
    prospetto = _prospetto(_apri(QUOTAZIONE(**QUOT_MONITORIO_10000_MINIMI))[1])
    etichette = [k for k in prospetto if "co. 1 bis" in k or "co. 1-bis" in k]
    assert etichette, prospetto
    assert all("fino al 30" in k for k in etichette), (
        f"etichetta {etichette!r}: art. 4, co. 1-bis, vigente prevede un aumento «fino al 30 per cento», "
        "non un aumento fisso del 30%"
    )


def test_quotazione_medi_coincidono_con_tabelle_gu(dm147_gu, out_dir):
    """Valori medi: tabella 8 (monitori) and tabella 2 (tribunale, used for the opposizione).

    Tabella 8 of D.M. 147/2022 has no scaglione up to 1.100 euro: the first one runs "da € 0 a
    € 5.200,00" (medio 473), which is what the tool applies to the plan case at 1.000 euro.
    """
    soglie8, fasi8 = _tabella_gu(dm147_gu, 8)
    assert soglie8 == [Decimal(s) for s in ("5200", "26000", "52000", "260000", "520000")]
    (medi8,) = fasi8.values()
    assert [(s, m) for s, _, m in pq._MONITORI_FASE_UNICA] == list(zip(soglie8, medi8))

    soglie2, fasi2 = _tabella_gu(dm147_gu, 2)
    assert soglie2 == [Decimal(s) for s in ("1100", "5200", "26000", "52000", "260000", "520000")]
    tabella_tool = pq._PARAMETRI["civile"]["scaglioni"]
    for chiave in ("studio", "introduttiva", "istruttoria", "decisionale"):
        gu = _fase(fasi2, chiave)
        tool = [Decimal(str(s[chiave]["medio"])) for s in tabella_tool[: len(gu)]]
        assert tool == gu, (chiave, tool, gu)
    assert [Decimal(str(s["fino_a"])) for s in tabella_tool[: len(soglie2)]] == soglie2

    # plan cases at medi: 473 (monitorio 1.000) and 919 + 777 + 1.680 + 1.701 (opposizione 26.000)
    assert _voce(_prospetto(_apri(QUOTAZIONE(**QUOT_MONITORIO_1000_MEDI))[1]), "Fase unica") == medi8[0]
    prospetto = _prospetto(_apri(QUOTAZIONE(**QUOT_OPPOSIZIONE_26000_MEDI))[1])
    assert [_voce(prospetto, f"Fase {f}") for f in ("di studio", "introduttiva", "istruttoria", "decisionale")] == [
        _fase(fasi2, k)[2] for k in ("studio", "introduttiva", "istruttoria", "decisionale")
    ]
    # scaglione boundaries of tabella 8: 5.200 is still the first, 5.200,01 the second
    for valore, atteso in ((5200, medi8[0]), (5200.01, medi8[1])):
        caso = {**QUOT_MONITORIO_1000_MEDI, "valore_causa": valore}
        assert _voce(_prospetto(_apri(QUOTAZIONE(**caso))[1]), "Fase unica") == atteso, valore


def test_quotazione_minimi_pari_alla_meta_esatta_del_medio(dm147_gu):
    """Valori minimi = medio ridotto del 50% (art. 4, co. 1, D.M. 55/2014), to the cent.

    The GU tables print only the medio. The tool (and src/data/parametri_forensi.json, "min")
    round the half UP to the euro when the medio is odd: 473 -> 237 (not 236,50), 567 -> 284
    (not 283,50), and the esecuzione defaults 331 -> 166, 567 -> 284. The rounded figure is
    above the legal minimum, so the letter's "valori minimi" overstate it by 0,50 per phase.
    """
    scarti = []
    _, fasi8 = _tabella_gu(dm147_gu, 8)
    (medi8,) = fasi8.values()
    for (soglia, minimo, _), medio in zip(pq._MONITORI_FASE_UNICA, medi8):
        if minimo != medio / 2:
            scarti.append(f"monitorio fino a {soglia}: tool {minimo}, norma {medio / 2}")

    _, fasi17 = _tabella_gu(dm147_gu, 17)
    for nome, default, medio in (
        ("introduttiva", pq._ESECUZIONE_DEFAULT_INTRODUTTIVA, _fase(fasi17, "introduttiva")[1]),
        ("trattazione e conclusiva", pq._ESECUZIONE_DEFAULT_TRATTAZIONE, _fase(fasi17, "trattazione")[1]),
    ):
        if Decimal(str(default)) != medio / 2:
            scarti.append(f"esecuzione default fase {nome}: tool {default}, norma {medio / 2}")

    _, fasi2 = _tabella_gu(dm147_gu, 2)
    for chiave in ("studio", "introduttiva", "istruttoria", "decisionale"):
        for scaglione, medio in zip(pq._PARAMETRI["civile"]["scaglioni"], _fase(fasi2, chiave)):
            minimo = Decimal(str(scaglione[chiave]["min"]))
            if minimo != medio / 2:
                scarti.append(f"opposizione fase {chiave} fino a {scaglione['fino_a']}: tool {minimo}, norma {medio / 2}")
    assert not scarti, "minimi arrotondati per eccesso rispetto al 50% esatto:\n" + "\n".join(scarti)


def test_quotazione_esecuzione_default_sono_tabella_17_presso_terzi(dm147_gu):
    """The esecuzione defaults 166/284 come from tabella 17 (presso terzi), not 16 (mobiliari).

    The code comment says "esecuzioni mobiliari, scaglione fino a € 5.200". Tabella 16 of
    D.M. 147/2022 (procedure esecutive mobiliari) has the phases studio + istruttoria/trattazione
    (1.100,01-5.200: medi 368 and 184, minimi 184 and 92); tabella 17 (presso terzi, consegna e
    rilascio) has fase introduttiva + fase di trattazione e conclusiva (medi 331 and 567,
    minimi 165,50 and 283,50): the tool's labels and figures (rounded to the euro) are those of
    tabella 17.
    """
    _, fasi16 = _tabella_gu(dm147_gu, 16)
    _, fasi17 = _tabella_gu(dm147_gu, 17)
    intro17, tratt17 = _fase(fasi17, "introduttiva"), _fase(fasi17, "trattazione")
    assert (intro17[1], tratt17[1]) == (Decimal("331"), Decimal("567"))
    assert abs(Decimal(str(pq._ESECUZIONE_DEFAULT_INTRODUTTIVA)) - intro17[1] / 2) <= Decimal("0.5")
    assert abs(Decimal(str(pq._ESECUZIONE_DEFAULT_TRATTAZIONE)) - tratt17[1] / 2) <= Decimal("0.5")
    studio16, tratt16 = _fase(fasi16, "studio"), _fase(fasi16, "trattazione")
    assert (studio16[1] / 2, tratt16[1] / 2) == (Decimal("184"), Decimal("92"))
    assert (Decimal("166"), Decimal("284")) != (studio16[1] / 2, tratt16[1] / 2)


def test_quotazione_esecuzione_default_sotto_1100_euro(dm147_gu, out_dir):
    """Esecuzione of 1.000 euro with the default compensi.

    The docstring says the 166/284 default "vale SOLO per valore causa fino a € 5.200 a livello
    minimi" and the guard only rejects values above 5.200. Tabella 17 has a first scaglione
    "da € 0,01 a € 1.100,00" (medi 110 and 236, minimi 55 and 118): below 1.100 euro the default
    overstates the tabellare by 277 euro (450 instead of 173).
    """
    soglie17, fasi17 = _tabella_gu(dm147_gu, 17)
    assert soglie17[0] == Decimal("1100")
    attesi = (_fase(fasi17, "introduttiva")[0] / 2, _fase(fasi17, "trattazione")[0] / 2)
    assert attesi == (Decimal("55"), Decimal("118"))

    caso = {**QUOT_ESECUZIONE_2500, "valore_causa": 1000}
    risultato = QUOTAZIONE(**caso)
    if risultato.startswith("Errore:"):
        return  # a refusal of the default below 1.100 euro is also a correct behaviour
    prospetto = _prospetto(_apri(risultato)[1])
    tool = (_voce(prospetto, "Fase introduttiva"), _voce(prospetto, "Fase di trattazione"))
    assert tool == attesi, (
        f"esecuzione 1.000 euro: il tool applica {tool}, la tabella 17 (scaglione fino a 1.100) "
        f"dà minimi {attesi}"
    )
