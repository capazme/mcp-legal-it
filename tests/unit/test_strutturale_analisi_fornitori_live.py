"""Structural live gate for `genera_report_fornitori` (src/tools/analisi_fornitori.py).

Benchmark phase 4, strategy "strutturale": the tool has no calculator on the
benchmark site, so the Excel it writes is checked against its own contract and
against the norm it declares (docstring ``Vigenza: art. 28 GDPR (nomina
responsabile); art. 4 GDPR (definizioni).``).

What it checks:

* the two cases of the benchmark plan (docs/benchmark/piano-benchmark-andreani.json):
  a valid batch (sheets, the 11 headers in their fixed order, row order with the
  no-DPA processors first, Avvertenze counters, the ``=HYPERLINK(...)`` ledger
  name stored as text) and a batch whose class carries two qualifications
  (rejected, no file written);
* edge cases: ``da_verificare`` sorted between ``no`` and ``si`` and not counted
  as a nomina; formula neutralisation on every third-party field, verified on
  the raw sheet XML (no ``<f>`` element); a traversal-shaped ``nome_file``;
  collect-all validation; control characters coming from a ledger extraction;
* the norm: art. 4 nn. 7-8 GDPR (definitions of titolare and responsabile, the
  source of the three qualifications) and art. 28 par. 3 and 9 GDPR (the
  processing by a processor is governed by a contract or other legal act in
  writing: why the no-DPA processors are the "nomine da predisporre"), read
  vigente through cite_law; the references declared by the tool are passed to
  verifica_citazioni.

The generated workbooks go to pytest's tmp_path (``_OUTPUT_DIR`` is patched),
never to the repository or to the user's temp folder. The date is pinned with
``LEGAL_TODAY``.

It needs the network (EUR-Lex/CELLAR) and is excluded from the default suite:

    .venv/bin/pytest tests/unit/test_strutturale_analisi_fornitori_live.py -m live -q -p no:cacheprovider -rfEs
"""

from __future__ import annotations

import asyncio
import json
import re
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

import src.server  # noqa: F401  (registers every tool module first: avoids a circular import)
from src.tools import analisi_fornitori
from tests.unit._norme_live import cite_law_json, contiene, normalizza

pytestmark = pytest.mark.live


def _fn(tool):
    return getattr(tool, "fn", tool)


REPORT = _fn(analisi_fornitori.genera_report_fornitori)

OGGI = "2026-09-25"

#: The 11 headers, in the order the tool promises (plan: "11 intestazioni nell'ordine fisso").
HEADER_ATTESO = [
    "Denominazione (da mastrino)",
    "P.IVA / CF",
    "Attività / servizi",
    "Categorie di dati presumibilmente trattate",
    "Qualificazione ipotizzata",
    "Motivazione sintetica",
    "Probabilità che tratti dati come responsabile",
    "DPA proprio del fornitore disponibile?",
    "Confidenza dell'identificazione",
    "Fonte (URL)",
    "Note / flag",
]

# ---------------------------------------------------------------------------
# Plan cases, copied verbatim from docs/benchmark/piano-benchmark-andreani.json
# ---------------------------------------------------------------------------

CASO_LOTTO_VALIDO = {
    "fornitori": [
        {"denominazione_mastrino": "Paghe Srl", "qualificazione": "responsabile", "motivazione": "elaborazione cedolini",
         "confidenza": "alto", "classe_attivita": "paghe", "probabilita_responsabile": "alta", "dpa_proprio": "no"},
        {"denominazione_mastrino": "Cloud Srl", "qualificazione": "responsabile", "motivazione": "hosting",
         "confidenza": "alto", "classe_attivita": "hosting", "probabilita_responsabile": "alta", "dpa_proprio": "si"},
        {"denominazione_mastrino": "Studio Rossi", "qualificazione": "titolare_autonomo", "motivazione": "difesa legale",
         "confidenza": "alto", "classe_attivita": "legale"},
        {"denominazione_mastrino": '=HYPERLINK("http://x")', "qualificazione": "fuori_perimetro",
         "motivazione": "fornitura carta", "confidenza": "medio", "classe_attivita": "cancelleria"},
    ],
    "cliente": "Alfa S.r.l.",
    "nome_file": "bench_report",
}

CASO_CLASSE_INCOERENTE = {
    "fornitori": [
        {"denominazione_mastrino": "Paghe Srl", "qualificazione": "responsabile", "motivazione": "elaborazione cedolini",
         "confidenza": "alto", "classe_attivita": "paghe", "probabilita_responsabile": "alta", "dpa_proprio": "no"},
        {"denominazione_mastrino": "Cloud Srl", "qualificazione": "responsabile", "motivazione": "hosting",
         "confidenza": "alto", "classe_attivita": "hosting", "probabilita_responsabile": "alta", "dpa_proprio": "si"},
        {"denominazione_mastrino": "Paghe2", "qualificazione": "titolare_autonomo", "motivazione": "x",
         "confidenza": "alto", "classe_attivita": "paghe"},
    ],
    "cliente": "Alfa S.r.l.",
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


@pytest.fixture
def out_dir(tmp_path, monkeypatch) -> Path:
    """Redirect the tool's output folder to tmp_path and pin today's date."""
    monkeypatch.setattr(analisi_fornitori, "_OUTPUT_DIR", str(tmp_path))
    monkeypatch.setenv("LEGAL_TODAY", OGGI)
    return tmp_path


def _resp(nome: str, dpa: str, classe: str = "hosting", **extra) -> dict:
    riga = {"denominazione_mastrino": nome, "qualificazione": "responsabile", "motivazione": "servizio per conto",
            "confidenza": "alto", "classe_attivita": classe, "probabilita_responsabile": "alta", "dpa_proprio": dpa}
    riga.update(extra)
    return riga


def _fuori(nome: str, classe: str = "cancelleria", **extra) -> dict:
    riga = {"denominazione_mastrino": nome, "qualificazione": "fuori_perimetro", "motivazione": "nessun dato personale",
            "confidenza": "medio", "classe_attivita": classe}
    riga.update(extra)
    return riga


def _file_salvato(out: str, out_dir: Path) -> Path:
    assert out.startswith("File salvato: "), out
    path = Path(out[len("File salvato: "):].rsplit(" (", 1)[0])
    assert path.parent == out_dir, f"il file deve finire nella cartella di output, non in {path.parent}"
    assert path.is_file(), path
    return path


def _avvertenze(path: Path) -> dict:
    ws = load_workbook(path)["Avvertenze"]
    return {row[0].value: row[1].value for row in ws.iter_rows() if row[0].value}


def _righe_analisi(path: Path) -> list[list]:
    ws = load_workbook(path)["Analisi fornitori"]
    return [[c.value for c in row] for row in ws.iter_rows(min_row=2)]


def _xml_fogli(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("xl/worksheets/sheet")}


def _nessuna_formula(path: Path) -> list[str]:
    """Names of the sheet parts that contain a formula element (must be empty)."""
    return [n for n, xml in _xml_fogli(path).items() if re.search(r"<f[\s>/]", xml)]


def _testo_vigente_o_skip(reference: str) -> str:
    payload = cite_law_json(reference)
    if payload.get("errore"):
        err = str(payload["errore"]).lower()
        if any(k in err for k in ("timeout", "connessione", "connection", "raggiungibile", "503", "502", "202")):
            pytest.skip(f"fonte temporaneamente non disponibile (sito_non_disponibile): {payload['errore']}")
        pytest.fail(f"cite_law({reference!r}) ha restituito un errore: {payload['errore']}")
    testo = payload.get("testo") or ""
    assert testo, payload
    return normalizza(testo)


# ---------------------------------------------------------------------------
# Plan case 1 -- "Lotto valido"
# ---------------------------------------------------------------------------


def test_caso_piano_lotto_valido_layout_ordinamento_art_28_gdpr(out_dir):
    """Fogli, 11 intestazioni, ordine (responsabile senza DPA in cima), contatori, formula come testo."""
    out = REPORT(**CASO_LOTTO_VALIDO)
    path = _file_salvato(out, out_dir)
    assert path.name == "bench_report.xlsx"

    wb = load_workbook(path)
    assert wb.sheetnames == ["Avvertenze", "Analisi fornitori"]
    ws = wb["Analisi fornitori"]
    assert [c.value for c in ws[1]] == HEADER_ATTESO
    assert ws.freeze_panes == "A2"

    righe = _righe_analisi(path)
    assert [r[0] for r in righe] == ["Paghe Srl", "Cloud Srl", "Studio Rossi", '=HYPERLINK("http://x")']
    assert [r[4] for r in righe] == [
        "Responsabile del trattamento", "Responsabile del trattamento",
        "Titolare autonomo", "Fuori perimetro privacy",
    ]
    # Probabilità e DPA solo per i responsabili, "—" altrimenti.
    assert [(r[6], r[7]) for r in righe] == [("Alta", "No"), ("Alta", "Sì"), ("—", "—"), ("—", "—")]
    assert [r[8] for r in righe] == ["Alto", "Alto", "Alto", "Medio"]

    # The ledger name '=HYPERLINK(...)' is third-party data: text, never a live formula.
    cella = ws.cell(row=5, column=1)
    assert cella.data_type != "f" and cella.value == '=HYPERLINK("http://x")'
    assert _nessuna_formula(path) == []

    avv = _avvertenze(path)
    assert avv["Cliente (titolare)"] == "Alfa S.r.l."
    assert avv["Data analisi"] == "25/09/2026", "default: oggi in formato gg/mm/aaaa"
    assert avv["File sorgente"] == "—"
    assert avv["Totale fornitori analizzati"] == 4
    assert avv["Responsabili del trattamento"] == 2
    assert avv["— di cui senza DPA proprio (nomina da predisporre)"] == 1
    assert avv["Titolari autonomi"] == 1
    assert avv["Fuori perimetro privacy"] == 1
    assert "da validare con il cliente e con i contratti" in avv["AVVERTENZE"]


# ---------------------------------------------------------------------------
# Plan case 2 -- "Classe con qualificazioni incoerenti"
# ---------------------------------------------------------------------------


def test_caso_piano_classe_incoerente_errore_e_nessun_file(out_dir):
    out = REPORT(**CASO_CLASSE_INCOERENTE)
    assert out.startswith("Errore di validazione"), out
    assert "classe 'paghe'" in out
    assert "responsabile: Paghe Srl" in out and "titolare_autonomo: Paghe2" in out
    assert "classe 'hosting'" not in out, "una classe coerente non va segnalata"
    assert list(out_dir.iterdir()) == []


# ---------------------------------------------------------------------------
# Edge cases
# ---------------------------------------------------------------------------


def test_da_verificare_fra_no_e_si_e_non_conta_come_nomina(out_dir):
    """Ordine: dpa no -> da_verificare -> si -> titolari -> fuori perimetro; A-Z senza distinzione di maiuscole."""
    fornitori = [
        _fuori("zeta carta"),
        _resp("beta cloud", "si"),
        _resp("Alfa Cloud", "da_verificare"),
        _resp("gamma cloud", "no"),
        _resp("Delta Cloud", "no"),
        {"denominazione_mastrino": "Studio Legale", "qualificazione": "titolare_autonomo",
         "motivazione": "mandato professionale", "confidenza": "alto", "classe_attivita": "legale"},
    ]
    path = _file_salvato(REPORT(fornitori=fornitori, cliente="Beta Spa", data_analisi="01/09/2026",
                                file_sorgente="mastrino.xlsx", nome_file="ordine"), out_dir)
    righe = _righe_analisi(path)
    assert [r[0] for r in righe] == ["Delta Cloud", "gamma cloud", "Alfa Cloud", "beta cloud", "Studio Legale",
                                     "zeta carta"]
    assert [r[7] for r in righe[:4]] == ["No", "No", "Da verificare", "Sì"]
    avv = _avvertenze(path)
    assert avv["Data analisi"] == "01/09/2026"
    assert avv["File sorgente"] == "mastrino.xlsx"
    assert avv["Responsabili del trattamento"] == 4
    assert avv["— di cui senza DPA proprio (nomina da predisporre)"] == 2


def test_neutralizzazione_formule_su_tutti_i_campi_di_terzi(out_dir):
    """Ogni campo di testo ('=' iniziale) resta testo: nessun elemento <f> nell'XML dei fogli."""
    payload = '=HYPERLINK("http://evil","clicca")'
    fornitori = [
        _resp(payload, "no", piva_cf="=1+1", attivita="=A1", categorie_dati="=B2", note="=C3",
              fonti=["=cmd|' /C calc'!A0", "https://esempio.it"]),
        _fuori("+SUM(1,1)", motivazione="@SUM(1)", note="-2+3"),
    ]
    path = _file_salvato(REPORT(fornitori=fornitori, cliente="=cliente()", file_sorgente="=sorgente()",
                                data_analisi="=oggi()", nome_file="formule"), out_dir)
    assert _nessuna_formula(path) == []
    righe = _righe_analisi(path)
    assert righe[0][0] == payload and righe[0][1] == "=1+1" and righe[0][10] == "=C3"
    assert righe[0][9] == "=cmd|' /C calc'!A0\nhttps://esempio.it"
    assert righe[1][0] == "+SUM(1,1)" and righe[1][5] == "@SUM(1)"
    avv = _avvertenze(path)
    assert avv["Cliente (titolare)"] == "=cliente()"
    assert avv["File sorgente"] == "=sorgente()"


def test_nome_file_con_percorso_resta_nella_cartella_di_output(out_dir):
    """OWASP A01 (path traversal): nome_file '../../x' non esce dalla cartella di output."""
    out = REPORT(fornitori=[_fuori("Carta Srl")], cliente="Gamma Srl", nome_file="../../fuga")
    path = _file_salvato(out, out_dir)
    assert path.name == "fuga.xlsx"
    assert not (out_dir.parent.parent / "fuga.xlsx").exists()


def test_validazione_collect_all_senza_file(out_dir):
    """Il contratto: tutti gli errori insieme, nessun file."""
    fornitori = [
        {"denominazione_mastrino": "Uno Srl", "qualificazione": "responsabile", "confidenza": "alto",
         "classe_attivita": "a", "probabilita_responsabile": "alta", "dpa_proprio": "no"},        # motivazione mancante
        _resp("Due Srl", "forse", classe="b"),                                                     # dpa non ammesso
        {"denominazione_mastrino": "Tre Srl", "qualificazione": "titolare_autonomo", "motivazione": "x",
         "confidenza": "alto", "classe_attivita": "c", "dpa_proprio": "no"},                       # dpa su non responsabile
        _fuori("Quattro Srl", classe="d", confidenza="certo"),                                     # confidenza non ammessa
    ]
    out = REPORT(fornitori=fornitori, cliente="Delta Srl")
    assert out.startswith("Errore di validazione"), out
    for frammento in ("riga 1: campo obbligatorio 'motivazione'", "riga 2: 'dpa_proprio' obbligatorio",
                      "riga 3: 'dpa_proprio' presente", "riga 4: 'confidenza' non valida"):
        assert frammento in out, (frammento, out)
    assert list(out_dir.iterdir()) == []


def test_caratteri_di_controllo_del_mastrino_errore_di_validazione(out_dir):
    """Denominazioni estratte da PDF/gestionali possono contenere caratteri di controllo (\\x0c salto pagina,
    \\x0b tab verticale) che XML/Excel non ammettono. Il contratto del tool ("Valida ogni riga e in caso di
    errori li restituisce tutti insieme senza scrivere il file") vuole un errore di validazione che segnali
    ENTRAMBE le righe (o una ripulitura), non un'eccezione di openpyxl che ne nomina una sola."""
    fornitori = [_fuori("ACME\x0cSRL", classe="varie"), _fuori("BETA\x0bSRL", classe="varie")]
    try:
        out = REPORT(fornitori=fornitori, cliente="Epsilon Srl", nome_file="controllo")
    except Exception as exc:  # noqa: BLE001 -- the failure mode under test
        pytest.fail(f"eccezione non gestita {type(exc).__name__}: {exc!r} (attesa: 'Errore di validazione' "
                    "su riga 1 e riga 2, oppure file con i caratteri di controllo rimossi)")
    if out.startswith("Errore di validazione"):
        assert "riga 1" in out and "riga 2" in out, out
        assert list(out_dir.iterdir()) == []
    else:
        path = _file_salvato(out, out_dir)
        assert [r[0] for r in _righe_analisi(path)] == ["ACMESRL", "BETASRL"]


# ---------------------------------------------------------------------------
# The norm: art. 4 nn. 7-8 and art. 28 GDPR, vigente text via cite_law
# ---------------------------------------------------------------------------


def test_art_4_nn_7_8_gdpr_definizioni_delle_qualificazioni():
    """Le tre qualificazioni del report poggiano sulle definizioni di titolare e responsabile (art. 4 nn. 7-8)."""
    testo = _testo_vigente_o_skip("art. 4 GDPR")
    missing = contiene(
        testo,
        "«titolare del trattamento»",
        "determina le finalità e i mezzi del trattamento di dati personali",
        "«responsabile del trattamento»",
        "che tratta dati personali per conto del titolare del trattamento",
    )
    assert not missing, f"art. 4 GDPR: il testo vigente non contiene {missing}"


def test_art_28_parr_3_e_9_gdpr_contratto_scritto_col_responsabile():
    """Perché i responsabili senza DPA proprio sono 'nomine da predisporre': art. 28 par. 3 e par. 9 GDPR."""
    testo = _testo_vigente_o_skip("art. 28 GDPR")
    missing = contiene(
        testo,
        "responsabile del trattamento",                       # rubrica = etichetta della colonna E
        "qualora un trattamento debba essere effettuato per conto del titolare del trattamento",
        "sono disciplinati da un contratto o da altro atto giuridico",
        "che vincoli il responsabile del trattamento al titolare del trattamento",
        "è stipulato in forma scritta",
    )
    assert not missing, f"art. 28 GDPR: il testo vigente non contiene {missing}"


def test_verifica_citazioni_riferimenti_dichiarati_dal_tool(out_dir):
    """Il report non cita norme nelle celle; i riferimenti sono quelli della riga Vigenza: del tool."""
    path = _file_salvato(REPORT(**CASO_LOTTO_VALIDO), out_dir)
    wb = load_workbook(path)
    celle = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
    assert re.findall(r"\bart\.\s*\d+", celle) == [], "nessun riferimento normativo atteso nelle celle"

    doc = analisi_fornitori.genera_report_fornitori.__doc__ or REPORT.__doc__ or ""
    if not doc:
        doc = getattr(analisi_fornitori.genera_report_fornitori, "description", "") or ""
    vigenza = next(line for line in doc.splitlines() if line.strip().startswith("Vigenza:"))
    riferimenti = [r.strip() for r in re.findall(r"art\.\s*\d+\s+GDPR", vigenza)]
    assert riferimenti == ["art. 28 GDPR", "art. 4 GDPR"], vigenza

    from src.tools.legal_citations import verifica_citazioni

    raw = asyncio.run(_fn(verifica_citazioni)(citazioni="\n".join(riferimenti), formato="json"))
    esito = json.loads(raw)
    assert esito.get("errore") is None, esito
    verdetti = {c["citazione"]: c["verdetto"] for c in esito["citazioni"]}
    if any(v == "non verificata" for v in verdetti.values()):
        pytest.skip(f"fonte temporaneamente non raggiungibile (sito_non_disponibile): {verdetti}")
    assert verdetti == {"art. 28 GDPR": "verificata", "art. 4 GDPR": "verificata"}, esito
